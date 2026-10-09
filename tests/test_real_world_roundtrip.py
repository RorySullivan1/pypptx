"""Round-trip and fidelity tests over the real-world corpus in `tests/test_files/real_world/`.

The corpus is PowerPoint-authored test data from Apache POI (see `SOURCES.md` there). Every
file must read cleanly through the public API, survive an edit, save and reload, and lose
nothing on a plain open-and-save.
"""

from __future__ import annotations

import gc
import io
import posixpath
import re
import sys
import zipfile
from pathlib import Path

import pytest
from lxml import etree

from pptx import Presentation
from pptx.exc import InvalidPackageError, PackageNotFoundError
from pptx.util import Inches

from .unitutil.deckwalk import walk

CORPUS_DIR = Path(__file__).parent / "test_files" / "real_world"
CORPUS = sorted(CORPUS_DIR.glob("*.pptx"), key=lambda p: p.name.lower())
MALFORMED = sorted((CORPUS_DIR / "malformed").glob("*.pptx"))
MARKER_TEXT = "pypptx round-trip marker"

# -- the corpus is test-only data kept out of published artifacts (MANIFEST.in prunes it from
# -- the sdist), so a test run from an unpacked sdist has no corpus directory at all. In a
# -- checkout the directory is tracked, and a missing deck fails `it_has_a_corpus_to_test`.
pytestmark = pytest.mark.skipif(
    not CORPUS_DIR.is_dir(), reason="real-world corpus not present (not shipped in the sdist)"
)

_blankless = etree.XMLParser(remove_blank_text=True)


def _canonical(xml: bytes) -> bytes:
    """`xml` in canonical form, ignoring whitespace-only text between elements."""
    return etree.tostring(etree.fromstring(xml, _blankless), method="c14n")


def _relationships(rels_xml: bytes) -> set[tuple[str | None, ...]]:
    """The relationships in a `.rels` part, as a set: their order carries no meaning."""
    return {
        (rel.get("Id"), rel.get("Type"), rel.get("Target"), rel.get("TargetMode", "Internal"))
        for rel in etree.fromstring(rels_xml)
    }


def _same_part(name: str, original: bytes, saved: bytes) -> bool:
    """True when the part `name` is unchanged.

    A `.rels` part compares as its set of relationships (pypptx writes them sorted by id), an
    XML part as canonical XML ignoring whitespace between elements, anything else byte for byte.
    XML that lxml can't parse or canonicalize (e.g. the SharePoint content-type schema in
    `bug68703.pptx`'s `customXml/item3.xml`) also compares byte for byte.
    """
    if name.endswith(".rels"):
        return _relationships(original) == _relationships(saved)
    if name.endswith(".xml"):
        try:
            return _canonical(original) == _canonical(saved)
        except (etree.XMLSyntaxError, etree.C14NError):
            pass
    return original == saved


def _content_types(zipf: zipfile.ZipFile) -> dict[str, str | None]:
    """Map each member's partname to the content type `[Content_Types].xml` gives it."""
    types = etree.fromstring(zipf.read("[Content_Types].xml"))
    ns = {"ct": "http://schemas.openxmlformats.org/package/2006/content-types"}
    defaults = {
        e.get("Extension").lower(): e.get("ContentType") for e in types.xpath("ct:Default", namespaces=ns)
    }
    overrides = {
        e.get("PartName").lower(): e.get("ContentType") for e in types.xpath("ct:Override", namespaces=ns)
    }
    mapping: dict[str, str | None] = {}
    for name in zipf.namelist():
        if name.endswith("/") or name == "[Content_Types].xml":
            continue
        partname = "/" + name
        ext = posixpath.basename(name).rpartition(".")[2].lower()
        mapping[partname] = overrides.get(partname.lower(), defaults.get(ext))
    return mapping


def it_has_a_corpus_to_test():
    assert len(CORPUS) >= 20, "expected the real-world corpus under %s" % CORPUS_DIR


@pytest.mark.realworld
@pytest.mark.parametrize("path", CORPUS, ids=[p.name for p in CORPUS])
class DescribeRealWorldDeck:
    """Each PowerPoint-authored deck in the corpus."""

    def it_reads_through_the_public_api_without_raising(self, path: Path):
        prs = Presentation(str(path))

        failures = walk(prs)

        assert failures == [], "\n".join("%s: %s" % f for f in failures[:20])

    def it_survives_an_edit_a_save_and_a_reload(self, path: Path):
        prs = Presentation(str(path))
        slide_count = len(prs.slides)
        if slide_count:
            textbox = prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
            textbox.text_frame.text = MARKER_TEXT
        stream = io.BytesIO()
        prs.save(stream)
        stream.seek(0)

        reloaded = Presentation(stream)

        assert len(reloaded.slides) == slide_count
        if slide_count:
            texts = [s.text_frame.text for s in reloaded.slides[0].shapes if s.has_text_frame]
            assert MARKER_TEXT in texts
        assert walk(reloaded) == []
        with zipfile.ZipFile(path) as original, zipfile.ZipFile(stream) as saved:
            missing = set(original.namelist()) - set(saved.namelist())
        assert missing == set()

    def it_changes_no_xml_on_a_plain_open_and_save(self, path: Path):
        stream = io.BytesIO()
        Presentation(str(path)).save(stream)

        with zipfile.ZipFile(path) as original, zipfile.ZipFile(stream) as saved:
            changed = [
                name
                for name in original.namelist()
                if not name.endswith("/")
                and name != "[Content_Types].xml"
                and not _same_part(name, original.read(name), saved.read(name))
            ]
            original_types = _content_types(original)
            saved_types = _content_types(saved)

        assert changed == []
        # -- every part keeps its content type, or gains one the original failed to declare --
        lost = {
            partname: (content_type, saved_types.get(partname))
            for partname, content_type in original_types.items()
            if content_type is not None and saved_types.get(partname) != content_type
        }
        assert lost == {}


def it_keeps_the_corpus_test_only():
    """The corpus is Apache POI test data, licensed for use here as test data only."""
    sources = (CORPUS_DIR / "SOURCES.md").read_text(encoding="utf-8")
    listed = set(re.findall(r"^\| `([^`]+\.pptx)` \| \d+ \| [^|]+ \| [^|]+ \|", sources, re.M))
    assert listed == {p.name for p in CORPUS}
    assert (CORPUS_DIR / "LICENSE").exists() and (CORPUS_DIR / "NOTICE").exists()


def _rebuilt(source: Path, drop: str) -> io.BytesIO:
    """A copy of the package at `source` without its `drop` member."""
    stream = io.BytesIO()
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            if info.filename != drop:
                dst.writestr(info, src.read(info.filename))
    stream.seek(0)
    return stream


@pytest.mark.realworld
class DescribeDamagedInput:
    """A file or stream that is not a readable package fails with a clear pypptx exception."""

    @pytest.mark.parametrize("path", MALFORMED, ids=[p.name[-24:] for p in MALFORMED])
    def it_rejects_each_malformed_corpus_file(self, path: Path):
        with pytest.raises(InvalidPackageError, match="is not a valid .pptx package"):
            Presentation(str(path))

    @pytest.mark.parametrize("kind", ["empty", "not-a-zip", "truncated"])
    def it_rejects_a_damaged_stream(self, kind: str):
        data = {
            "empty": lambda: b"",
            "not-a-zip": lambda: b"not a zip at all",
            "truncated": lambda: (CORPUS_DIR / "bar-chart.pptx").read_bytes()[:5000],
        }[kind]()

        with pytest.raises(InvalidPackageError, match="the stream is not a valid .pptx package"):
            Presentation(io.BytesIO(data))

    @pytest.mark.parametrize(
        ("drop", "message"),
        [
            ("[Content_Types].xml", r"has no \[Content_Types\]\.xml"),
            ("_rels/.rels", "has no main document part"),
            ("ppt/presentation.xml", "has no main document part"),
        ],
    )
    def it_rejects_a_zip_missing_a_required_part(self, drop: str, message: str):
        stream = _rebuilt(CORPUS_DIR / "bar-chart.pptx", drop)

        with pytest.raises(InvalidPackageError, match=message):
            Presentation(stream)

    def it_does_not_disguise_a_closed_stream_as_a_damaged_package(self):
        # -- a caller bug, not a bad file: the ValueError from the closed stream propagates --
        stream = io.BytesIO((CORPUS_DIR / "bar-chart.pptx").read_bytes())
        stream.close()

        with pytest.raises(ValueError, match="closed file") as e:
            Presentation(stream)
        assert not isinstance(e.value, InvalidPackageError)

    def it_still_reports_a_missing_path_as_not_found(self, tmp_path: Path):
        with pytest.raises(PackageNotFoundError):
            Presentation(str(tmp_path / "absent.pptx"))

    def it_leaves_no_error_behind_for_garbage_collection(self, monkeypatch):
        """Rejecting a file must not leave a half-built reader whose `__del__` raises later."""
        unraisable = []
        monkeypatch.setattr(sys, "unraisablehook", unraisable.append)

        for path in MALFORMED:
            with pytest.raises(InvalidPackageError):
                Presentation(str(path))
        gc.collect()

        assert [str(u.exc_value) for u in unraisable] == []


@pytest.mark.realworld
def it_keeps_in_document_hyperlinks_internal_when_a_slide_is_duplicated():
    """`with_japanese.pptx` links to footnote anchors ("#_ftn1"), which are internal targets."""
    prs = Presentation(str(CORPUS_DIR / "with_japanese.pptx"))
    duplicate = prs.slides.duplicate(prs.slides[0])
    stream = io.BytesIO()
    prs.save(stream)

    with zipfile.ZipFile(stream) as saved:
        slide_xml = saved.read(duplicate.part.partname.membername)
        rels_xml = saved.read("ppt/slides/_rels/%s.rels" % duplicate.part.partname.filename)
    rels = etree.fromstring(rels_xml)
    fragments = [rel for rel in rels if rel.get("Target").startswith("#")]

    # -- the duplicate relates each distinct target once; both anchors survive, as internal --
    assert {rel.get("Target") for rel in fragments} == {"#_ftn1", "#_ftnref1"}
    assert {rel.get("TargetMode", "Internal") for rel in fragments} == {"Internal"}
    # -- and every relationship the duplicated slide refers to exists --
    referenced = set(re.findall(rb'r:(?:id|embed|link)="(rId\d+)"', slide_xml))
    assert referenced <= {rel.get("Id").encode() for rel in rels}
