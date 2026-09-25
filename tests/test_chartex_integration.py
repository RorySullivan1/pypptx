"""Integration tests for chartex charts: reading PowerPoint-style files, adding charts, and
shape-tree operations on the `mc:AlternateContent` wrapper a chartex chart lives in."""

from __future__ import annotations

import pytest

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.oxml.shapes.shared import tree_elm
from pptx.parts.chartex import ChartExPart
from pptx.shapes.graphfrm import GraphicFrame
from pptx.util import Inches

from .unitutil.chartex import (
    HISTOGRAM_VALUES,
    TREEMAP_PATHS,
    TREEMAP_VALUES,
    WATERFALL_CATEGORIES,
    WATERFALL_VALUES,
    chartex_pptx,
    xlsx_column,
)
from .unitutil.cxml import element


def _waterfall_data() -> CategoryChartData:
    chart_data = CategoryChartData()
    chart_data.categories = ["Start", "Q1", "Q2", "End"]
    chart_data.add_series("Flow", (100, -20, 35, 115))
    return chart_data


def _slide_with_chartex():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(1))
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.WATERFALL, Inches(1), Inches(1), Inches(4), Inches(3), _waterfall_data()
    )
    slide.shapes.add_textbox(Inches(2), Inches(2), Inches(1), Inches(1))
    return slide, frame


class DescribeReadingChartexCharts:
    """Chartex charts in a file laid out the way PowerPoint saves them."""

    def it_lists_chartex_frames_among_the_slide_shapes(self):
        shapes = Presentation(chartex_pptx()).slides[0].shapes

        assert len(shapes) == 3
        assert [shape.shape_type for shape in shapes] == [MSO_SHAPE_TYPE.CHART] * 3
        assert all(shape.has_chartex and not shape.has_chart for shape in shapes)
        assert all(isinstance(shape.chartex_part, ChartExPart) for shape in shapes)
        assert [shape.chartex.chart_type for shape in shapes] == [
            XL_CHART_TYPE.WATERFALL,
            XL_CHART_TYPE.PARETO,
            XL_CHART_TYPE.TREEMAP,
        ]

    def it_reads_values_that_match_the_source_workbook(self):
        waterfall, pareto, treemap = Presentation(chartex_pptx()).slides[0].shapes

        def workbook(shape):
            return shape.chartex_part.xlsx_part.blob

        (series,) = waterfall.chartex.series
        assert series.name == "Flow"
        assert list(series.categories) == WATERFALL_CATEGORIES == xlsx_column(workbook(waterfall), "A")
        assert list(series.values) == WATERFALL_VALUES == xlsx_column(workbook(waterfall), "B")

        histogram, pareto_line = pareto.chartex.series
        assert list(histogram.values) == HISTOGRAM_VALUES == xlsx_column(workbook(pareto), "A")
        assert (pareto_line.layout_id, pareto_line.values) == ("paretoLine", ())

        (series,) = treemap.chartex.series
        assert list(series.category_paths) == TREEMAP_PATHS
        assert list(series.categories) == xlsx_column(workbook(treemap), "B")
        assert list(series.values) == TREEMAP_VALUES == xlsx_column(workbook(treemap), "C")

    def it_ignores_alternate_content_that_is_not_a_chartex_chart(self):
        spTree = element(
            "p:spTree/(p:sp,mc:AlternateContent/mc:Choice/p:sp,"
            "mc:AlternateContent/mc:Choice/p:graphicFrame/a:graphic/a:graphicData{uri=foo})"
        )

        assert [elm.tag for elm in spTree.iter_shape_elms()] == [qn("p:sp")]


class DescribeAddingChartexCharts:
    """`shapes.add_chart()` with a chartex chart type."""

    @pytest.mark.parametrize(
        "chart_type",
        [
            XL_CHART_TYPE.WATERFALL,
            XL_CHART_TYPE.HISTOGRAM,
            XL_CHART_TYPE.BOX_WHISKER,
            XL_CHART_TYPE.TREEMAP,
            XL_CHART_TYPE.SUNBURST,
            XL_CHART_TYPE.FUNNEL,
        ],
    )
    def it_adds_a_graphic_frame_wrapped_in_alternate_content(self, chart_type):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        frame = slide.shapes.add_chart(
            chart_type, Inches(1), Inches(1), Inches(4), Inches(3), _waterfall_data()
        )

        assert frame.has_chartex
        assert frame.chartex.chart_type is chart_type
        assert (frame.left, frame.top, frame.width, frame.height) == (
            Inches(1),
            Inches(1),
            Inches(4),
            Inches(3),
        )
        alternateContent = tree_elm(frame.element)
        assert alternateContent.tag == qn("mc:AlternateContent")
        choice, fallback = alternateContent
        assert choice.get("Requires") == "cx1"
        assert fallback.xpath("./p:sp/p:nvSpPr/p:cNvPr/@id") == [str(frame.shape_id)]
        assert list(slide.shapes) == [frame]

    def but_it_rejects_chart_data_the_chart_type_cannot_show(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        chart_data = CategoryChartData()
        chart_data.add_series("Sales", (1, 2))

        with pytest.raises(ValueError, match="needs categories"):
            slide.shapes.add_chart(
                XL_CHART_TYPE.TREEMAP, Inches(1), Inches(1), Inches(4), Inches(3), chart_data
            )


class DescribeChartexShapeTreeOperations:
    """Moving, copying and removing a chartex chart acts on its whole alternate-content wrapper."""

    def it_moves_the_whole_wrapper_to_the_front_and_back(self):
        slide, frame = _slide_with_chartex()
        shapes = slide.shapes

        shapes.move_shape_to_front(frame)
        assert shapes.index(frame) == 2
        assert tree_elm(frame.element).getparent() is shapes._spTree

        shapes.move_shape_to_back(frame)
        assert shapes.index(frame) == 0
        assert frame.element.getparent().tag == qn("mc:Choice")

    def it_moves_another_shape_behind_the_wrapper(self):
        slide, frame = _slide_with_chartex()
        shapes = slide.shapes
        shapes.move_shape_to_back(frame)
        textbox = shapes[2]

        shapes.move_shape_to_back(textbox)

        assert shapes.index(textbox) == 0
        assert shapes.index(frame) == 1

    def it_removes_the_whole_wrapper(self):
        slide, frame = _slide_with_chartex()

        slide.shapes.remove_shape(frame)

        assert len(slide.shapes) == 2
        assert slide.shapes._spTree.xpath("./mc:AlternateContent") == []

    def it_duplicates_the_wrapper_with_a_new_id_on_the_frame_and_its_fallback(self):
        slide, frame = _slide_with_chartex()

        copy = slide.shapes.duplicate_shape(frame)

        assert copy.has_chartex
        assert copy.shape_id != frame.shape_id
        wrapper = tree_elm(copy.element)
        assert wrapper is not tree_elm(frame.element)
        assert set(wrapper.xpath(".//p:cNvPr/@id")) == {str(copy.shape_id)}
        assert set(wrapper.xpath(".//p:cNvPr/@name")) == {copy.name}
        assert len(slide.shapes) == 4

    def it_groups_the_whole_wrapper(self):
        slide, frame = _slide_with_chartex()

        group = slide.shapes.add_group_shape([frame])

        (grouped,) = group.shapes
        assert grouped.has_chartex
        assert grouped.is_in_group is True
        assert tree_elm(grouped.element).getparent() is group.element
        assert len(slide.shapes) == 3
        assert (group.left, group.top, group.width, group.height) == (
            Inches(1),
            Inches(1),
            Inches(4),
            Inches(3),
        )

    def it_knows_a_top_level_chartex_chart_is_not_in_a_group(self):
        _, frame = _slide_with_chartex()

        assert frame.is_in_group is False

    def it_finds_the_group_holding_a_chartex_chart(self):
        grpSp = element(
            "p:grpSp/mc:AlternateContent/mc:Choice/p:graphicFrame/a:graphic/"
            "a:graphicData{uri=http://schemas.microsoft.com/office/drawing/2014/chartex}"
        )
        frame = GraphicFrame(grpSp.xpath(".//p:graphicFrame")[0], None)

        assert frame.is_in_group is True
        assert frame.parent_group.element is grpSp
