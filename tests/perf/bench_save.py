"""Save-path benchmark.

Times `Presentation.save(BytesIO)` after one trivial mutation. Isolates the
serialize-and-zip cost from cold-open cost.
"""

from __future__ import annotations

import io
import time
from pathlib import Path

import pytest

from pptx import Presentation
from pptx.util import Inches


def _time_save(path: Path) -> float:
    prs = Presentation(str(path))
    # Trivial mutation to force XML write paths.
    slide = prs.slides[0]
    slide.shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(1.0), Inches(0.4))
    buf = io.BytesIO()
    t0 = time.perf_counter()
    prs.save(buf)
    return time.perf_counter() - t0


@pytest.mark.perf
def bench_save_small(synth_small: Path) -> None:
    elapsed = _time_save(synth_small)
    print(f"\n[bench_save small] {elapsed * 1000:.2f} ms")


@pytest.mark.perf
def bench_save_medium(synth_medium: Path) -> None:
    elapsed = _time_save(synth_medium)
    print(f"\n[bench_save medium] {elapsed * 1000:.2f} ms")


@pytest.mark.perf
def bench_save_large(synth_large: Path) -> None:
    elapsed = _time_save(synth_large)
    print(f"\n[bench_save large] {elapsed * 1000:.2f} ms")
