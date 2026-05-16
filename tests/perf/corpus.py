"""Synthetic corpus generator for the performance harness.

Builds `.pptx` files programmatically via the public API so the harness has
deterministic, parametrized fixtures to time. Modeled loosely on
`examples/investor_presentation.py`.

Used by `tests/perf/conftest.py` (pytest fixtures) and
`tests/perf/run_baseline.py` (the standalone baseline runner).
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

TEST_FILES = Path(__file__).resolve().parent.parent / "test_files"
IMAGE_PATH = TEST_FILES / "python-powered.png"


@dataclass(frozen=True)
class CorpusSpec:
    """Parameters that fully describe a synthetic deck."""

    name: str
    n_slides: int
    shapes_per_slide: int
    paragraphs_per_shape: int
    n_images: int = 0


SMALL = CorpusSpec(
    name="small", n_slides=5, shapes_per_slide=5, paragraphs_per_shape=2
)
MEDIUM = CorpusSpec(
    name="medium",
    n_slides=30,
    shapes_per_slide=15,
    paragraphs_per_shape=4,
    n_images=5,
)
LARGE = CorpusSpec(
    name="large",
    n_slides=100,
    shapes_per_slide=25,
    paragraphs_per_shape=5,
    n_images=10,
)

# Single-slide spec used for the index-linearity benchmark.
LOOKUP_COUNTS = (10, 100, 1000)


def build_deck(spec: CorpusSpec, dest: Path) -> Path:
    """Write a `.pptx` matching `spec` to `dest`. Returns `dest`."""
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    image_path_str = str(IMAGE_PATH)
    image_budget = spec.n_images

    for slide_idx in range(spec.n_slides):
        slide = prs.slides.add_slide(blank_layout)
        for shape_idx in range(spec.shapes_per_slide):
            tb = slide.shapes.add_textbox(
                Inches(0.5),
                Inches(0.5 + shape_idx * 0.1),
                Inches(4.0),
                Inches(0.5),
            )
            tf = tb.text_frame
            tf.text = f"slide={slide_idx} shape={shape_idx} p=0"
            tf.paragraphs[0].font.size = Pt(12)
            for p_idx in range(1, spec.paragraphs_per_shape):
                p = tf.add_paragraph()
                p.text = f"slide={slide_idx} shape={shape_idx} p={p_idx}"
                p.font.size = Pt(11)
        if image_budget > 0:
            slide.shapes.add_picture(
                image_path_str, Inches(5.0), Inches(0.5), Inches(2.0), Inches(2.0)
            )
            image_budget -= 1

    prs.save(str(dest))
    return dest


def build_lookup_deck(n_shapes: int, dest: Path) -> Path:
    """Single-slide deck with `n_shapes` textboxes for index-lookup timing."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i in range(n_shapes):
        tb = slide.shapes.add_textbox(
            Inches(0.1),
            Inches(0.1 + (i % 50) * 0.1),
            Inches(0.5),
            Inches(0.2),
        )
        tb.name = f"box_{i}"
    prs.save(str(dest))
    return dest
