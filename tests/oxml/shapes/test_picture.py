"""Unit-test suite for `pptx.oxml.shapes.picture` module."""

from __future__ import annotations

from typing import cast

import pytest

from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.oxml.shapes.picture import P14_MEDIA_EXT_URI, CT_Picture

from ...unitutil.cxml import element


class DescribeCT_Picture:
    """Unit-test suite for `pptx.oxml.shapes.picture.CT_Picture` objects."""

    @pytest.mark.parametrize(
        "desc, xml_desc",
        (
            ("kittens.jpg", "kittens.jpg"),
            ("bits&bobs.png", "bits&amp;bobs.png"),
            ("img&.png", "img&amp;.png"),
            ("im<ag>e.png", "im&lt;ag&gt;e.png"),
        ),
    )
    def it_can_create_a_new_pic_element(self, desc, xml_desc):
        """`desc` attr (often filename) is XML-escaped to handle special characters.

        In particular, ampersand ('&'), less/greater-than ('</>') etc.
        """
        pic = CT_Picture.new_pic(
            shape_id=9, name="Picture 8", desc=desc, rId="rId42", x=1, y=2, cx=3, cy=4
        )

        assert pic.xml == (
            "<p:pic %s>\n"
            "  <p:nvPicPr>\n"
            '    <p:cNvPr id="9" name="Picture 8" descr="%s"/>\n'
            "    <p:cNvPicPr>\n"
            '      <a:picLocks noChangeAspect="1"/>\n'
            "    </p:cNvPicPr>\n"
            "    <p:nvPr/>\n"
            "  </p:nvPicPr>\n"
            "  <p:blipFill>\n"
            '    <a:blip r:embed="rId42"/>\n'
            "    <a:stretch>\n"
            "      <a:fillRect/>\n"
            "    </a:stretch>\n"
            "  </p:blipFill>\n"
            "  <p:spPr>\n"
            "    <a:xfrm>\n"
            '      <a:off x="1" y="2"/>\n'
            '      <a:ext cx="3" cy="4"/>\n'
            "    </a:xfrm>\n"
            '    <a:prstGeom prst="rect">\n'
            "      <a:avLst/>\n"
            "    </a:prstGeom>\n"
            "  </p:spPr>\n"
            "</p:pic>\n" % (nsdecls("a", "p", "r"), xml_desc)
        )

    def it_can_create_a_new_video_pic_element(self):
        pic = CT_Picture.new_video_pic(
            shape_id=42,
            shape_name="Media 41",
            video_rId="rId1",
            media_rId="rId2",
            poster_frame_rId="rId3",
            x=1,
            y=2,
            cx=3,
            cy=4,
        )

        assert pic.xml == (
            "<p:pic %s>\n"
            "  <p:nvPicPr>\n"
            '    <p:cNvPr id="42" name="Media 41">\n'
            '      <a:hlinkClick r:id="" action="ppaction://media"/>\n'
            "    </p:cNvPr>\n"
            "    <p:cNvPicPr>\n"
            '      <a:picLocks noChangeAspect="1"/>\n'
            "    </p:cNvPicPr>\n"
            "    <p:nvPr>\n"
            '      <a:videoFile r:link="rId1"/>\n'
            "      <p:extLst>\n"
            '        <p:ext uri="{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}">\n'
            '          <p14:media xmlns:p14="http://schemas.microsoft.com/office/power'
            'point/2010/main" r:embed="rId2"/>\n'
            "        </p:ext>\n"
            "      </p:extLst>\n"
            "    </p:nvPr>\n"
            "  </p:nvPicPr>\n"
            "  <p:blipFill>\n"
            '    <a:blip r:embed="rId3"/>\n'
            "    <a:stretch>\n"
            "      <a:fillRect/>\n"
            "    </a:stretch>\n"
            "  </p:blipFill>\n"
            "  <p:spPr>\n"
            "    <a:xfrm>\n"
            '      <a:off x="1" y="2"/>\n'
            '      <a:ext cx="3" cy="4"/>\n'
            "    </a:xfrm>\n"
            '    <a:prstGeom prst="rect">\n'
            "      <a:avLst/>\n"
            "    </a:prstGeom>\n"
            "  </p:spPr>\n"
            "</p:pic>\n"
        ) % nsdecls("a", "p", "r")

    def it_can_create_a_new_audio_pic_element(self):
        pic = CT_Picture.new_audio_pic(
            shape_id=42,
            shape_name="Audio 41",
            audio_rId="rId1",
            media_rId="rId2",
            icon_rId="rId3",
            x=1,
            y=2,
            cx=3,
            cy=4,
        )

        assert pic.xml == (
            "<p:pic %s>\n"
            "  <p:nvPicPr>\n"
            '    <p:cNvPr id="42" name="Audio 41">\n'
            '      <a:hlinkClick r:id="" action="ppaction://media"/>\n'
            "    </p:cNvPr>\n"
            "    <p:cNvPicPr>\n"
            '      <a:picLocks noChangeAspect="1"/>\n'
            "    </p:cNvPicPr>\n"
            "    <p:nvPr>\n"
            '      <a:audioFile r:link="rId1"/>\n'
            "      <p:extLst>\n"
            '        <p:ext uri="{DAA4B4D4-6D71-4841-9C94-3DE7FCFB9230}">\n'
            '          <p14:media xmlns:p14="http://schemas.microsoft.com/office/power'
            'point/2010/main" r:embed="rId2"/>\n'
            "        </p:ext>\n"
            "      </p:extLst>\n"
            "    </p:nvPr>\n"
            "  </p:nvPicPr>\n"
            "  <p:blipFill>\n"
            '    <a:blip r:embed="rId3"/>\n'
            "    <a:stretch>\n"
            "      <a:fillRect/>\n"
            "    </a:stretch>\n"
            "  </p:blipFill>\n"
            "  <p:spPr>\n"
            "    <a:xfrm>\n"
            '      <a:off x="1" y="2"/>\n'
            '      <a:ext cx="3" cy="4"/>\n'
            "    </a:xfrm>\n"
            '    <a:prstGeom prst="rect">\n'
            "      <a:avLst/>\n"
            "    </a:prstGeom>\n"
            "  </p:spPr>\n"
            "</p:pic>\n"
        ) % nsdecls("a", "p", "r")

    def it_has_no_p14_media_when_none_is_present(self):
        pic = element("p:pic/p:nvPicPr/p:nvPr")
        assert pic.p14_media is None

    def it_provides_access_to_its_p14_media_when_present(self):
        pic = cast(
            "CT_Picture",
            parse_xml(
                "<p:pic %s>\n"
                "  <p:nvPicPr>\n"
                "    <p:nvPr>\n"
                "      <p:extLst>\n"
                '        <p:ext uri="%s">\n'
                '          <p14:media r:embed="rId9"/>\n'
                "        </p:ext>\n"
                "      </p:extLst>\n"
                "    </p:nvPr>\n"
                "  </p:nvPicPr>\n"
                "</p:pic>" % (nsdecls("p", "p14", "r"), P14_MEDIA_EXT_URI)
            ),
        )
        p14_media = pic.p14_media
        assert p14_media is not None
        assert p14_media.embed == "rId9"

    def it_adds_the_extLst_wrapper_when_getting_or_adding_p14_media(self):
        pic = element("p:pic/(p:nvPicPr/p:nvPr,p:blipFill,p:spPr)")

        media = pic.get_or_add_p14_media()

        assert media is pic.p14_media
        ext = pic.xpath("./p:nvPicPr/p:nvPr/p:extLst/p:ext")[0]
        assert ext.get("uri") == P14_MEDIA_EXT_URI


class DescribeCT_Media:
    """Unit-test suite for `pptx.oxml.shapes.picture.CT_Media` objects."""

    def it_can_get_and_set_r_embed(self):
        media = element("p14:media")
        assert media.embed is None
        media.embed = "rId4"
        assert media.embed == "rId4"
        media.embed = None
        assert media.embed is None

    def it_can_get_and_set_its_trim(self):
        media = element("p14:media")
        assert media.trim is None

        trim = media.get_or_add_trim()
        trim.st = 1500
        trim.end = 2000

        assert media.trim.st == 1500
        assert media.trim.end == 2000

    def it_can_get_and_set_its_fade(self):
        media = element("p14:media")
        assert media.fade is None

        fade = media.get_or_add_fade()
        fade.in_ = 500
        fade.out = 750

        assert media.fade.in_ == 500
        assert media.fade.out == 750
