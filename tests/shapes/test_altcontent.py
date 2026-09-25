"""Unit-test suite for `pptx.shapes.altcontent` module."""

from __future__ import annotations

import pytest

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import ShapeError
from pptx.oxml.ns import qn
from pptx.shapes.altcontent import AlternateContentShape

from ..unitutil.cxml import element

MODEL3D_URI = "http://schemas.microsoft.com/office/drawing/2017/model3d"
ZOOM_URI = "http://schemas.microsoft.com/office/powerpoint/2016/sectionzoom"


def _shape(alternateContent_cxml: str, container: str = "mc:Fallback") -> AlternateContentShape:
    alternateContent = element(alternateContent_cxml)
    shape_elm = alternateContent.xpath("./%s/*" % container)[0]
    return AlternateContentShape(shape_elm, None)


class DescribeAlternateContentShape:
    """Unit-test suite for `pptx.shapes.altcontent.AlternateContentShape` objects."""

    @pytest.mark.parametrize(
        ("choice_cxml", "expected_value"),
        [
            (
                "p:graphicFrame/a:graphic/a:graphicData{uri=%s}" % MODEL3D_URI,
                "model3d",
            ),
            ("p:graphicFrame/a:graphic/a:graphicData{uri=%s}" % ZOOM_URI, "zoom"),
            ("p:sp/p:txBody/a:p/a14:m/m:oMathPara/m:oMath", "equation"),
            ("p:sp/p:txBody/a:p/a:r/a:t", "unknown"),
            ("p:graphicFrame/a:graphic/a:graphicData{uri=foo}", "unknown"),
        ],
    )
    def it_knows_what_kind_of_content_it_holds(self, choice_cxml: str, expected_value: str):
        shape = _shape("mc:AlternateContent/(mc:Choice/%s,mc:Fallback/p:pic)" % choice_cxml)

        assert shape.content_kind == expected_value

    def it_recognizes_content_from_the_namespace_its_choice_requires(self):
        alternateContent = element("mc:AlternateContent/(mc:Choice,mc:Fallback/p:pic)")
        choice = alternateContent[0]
        # -- cxml can't declare an arbitrary prefix, so re-parent the choice under one that does --
        from pptx.oxml import parse_xml

        declared = parse_xml(
            '<mc:Choice xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
            ' xmlns:am3d="%s" Requires="am3d"/>' % MODEL3D_URI
        )
        alternateContent.replace(choice, declared)
        shape = AlternateContentShape(alternateContent.xpath("./mc:Fallback/p:pic")[0], None)

        assert shape.content_kind == "model3d"
        assert shape.shape_type is MSO_SHAPE_TYPE.MODEL_3D

    @pytest.mark.parametrize(
        ("cxml", "container", "expected_value"),
        [
            ("mc:AlternateContent/(mc:Choice/p:sp,mc:Fallback/p:pic)", "mc:Fallback", True),
            ("mc:AlternateContent/(mc:Choice/p:sp,mc:Fallback)", "mc:Choice", False),
            ("mc:AlternateContent/mc:Choice/p:sp", "mc:Choice", False),
        ],
    )
    def it_knows_whether_it_has_a_fallback(self, cxml: str, container: str, expected_value: bool):
        assert _shape(cxml, container).has_fallback is expected_value

    def it_provides_access_to_its_wrapper_element(self):
        shape = _shape("mc:AlternateContent/(mc:Choice/p:sp,mc:Fallback/p:pic)")

        assert shape.alternate_content_element.tag == qn("mc:AlternateContent")

    def it_has_no_shape_type_unless_it_is_a_3d_model(self):
        shape = _shape(
            "mc:AlternateContent/(mc:Choice/p:graphicFrame/a:graphic/a:graphicData{uri=%s},"
            "mc:Fallback/p:pic)" % ZOOM_URI
        )

        assert shape.shape_type is None

    def it_reads_its_name_and_geometry_from_the_shape_it_stands_for(self):
        shape = _shape(
            "mc:AlternateContent/(mc:Choice/p:sp,mc:Fallback/p:pic/("
            "p:nvPicPr/p:cNvPr{id=7,name=Zoom},"
            "p:spPr/a:xfrm/(a:off{x=10,y=20},a:ext{cx=30,cy=40})))"
        )

        assert (shape.shape_id, shape.name) == (7, "Zoom")
        assert (shape.left, shape.top, shape.width, shape.height) == (10, 20, 30, 40)

    @pytest.mark.parametrize(
        ("prop", "value"),
        [
            ("left", 1),
            ("top", 1),
            ("width", 1),
            ("height", 1),
            ("rotation", 45.0),
            ("name", "foo"),
        ],
    )
    def it_refuses_changes_to_its_geometry_and_name(self, prop: str, value: object):
        shape = _shape(
            "mc:AlternateContent/(mc:Choice/p:sp,mc:Fallback/p:pic/p:nvPicPr/p:cNvPr{id=7,name=Z})"
        )

        with pytest.raises(ShapeError, match="read-only"):
            setattr(shape, prop, value)
