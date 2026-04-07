# pyright: reportPrivateUsage=false

"""Unit test suite for pptx.shapes.range module."""

from __future__ import annotations

import pytest

from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.shapes.base import BaseShape
from pptx.shapes.range import ShapeRange
from pptx.util import Emu, Inches


class FakeParent:
    """Minimal stand-in for ProvidesPart parent used by BaseShape."""

    @property
    def part(self):
        return None


def _make_shape(left: int, top: int, width: int, height: int) -> BaseShape:
    """Return a BaseShape with the given position and size (in EMU)."""
    xml = (
        '<p:sp %s>'
        '  <p:nvSpPr>'
        '    <p:cNvPr id="1" name="shape"/>'
        '    <p:cNvSpPr/>'
        '    <p:nvPr/>'
        '  </p:nvSpPr>'
        '  <p:spPr>'
        '    <a:xfrm>'
        '      <a:off x="%d" y="%d"/>'
        '      <a:ext cx="%d" cy="%d"/>'
        '    </a:xfrm>'
        '  </p:spPr>'
        '</p:sp>' % (nsdecls("a", "p", "r"), left, top, width, height)
    )
    elm = parse_xml(xml)
    return BaseShape(elm, FakeParent())


class DescribeShapeRange:
    """Unit tests for ShapeRange."""

    def it_rejects_empty_input(self):
        with pytest.raises(ValueError, match="at least one shape"):
            ShapeRange([])

    def it_provides_len_iter_and_getitem(self):
        shapes = [_make_shape(0, 0, 100, 100), _make_shape(200, 0, 100, 100)]
        sr = ShapeRange(shapes)
        assert len(sr) == 2
        assert list(sr) == shapes
        assert sr[0] is shapes[0]
        assert sr[1] is shapes[1]

    def it_computes_bounding_box(self):
        s1 = _make_shape(100, 200, 300, 400)
        s2 = _make_shape(500, 100, 200, 600)
        sr = ShapeRange([s1, s2])
        assert sr.bbox_left == 100
        assert sr.bbox_top == 100
        assert sr.bbox_right == 700  # 500 + 200
        assert sr.bbox_bottom == 700  # 100 + 600
        assert sr.bbox_width == 600  # 700 - 100
        assert sr.bbox_height == 600  # 700 - 100


class DescribeShapeRange_Alignment:
    """Unit tests for ShapeRange alignment methods."""

    def it_aligns_left(self):
        s1 = _make_shape(100, 0, 50, 50)
        s2 = _make_shape(300, 0, 80, 50)
        sr = ShapeRange([s1, s2])
        sr.align_left()
        assert s1.left == 100
        assert s2.left == 100

    def it_aligns_center(self):
        s1 = _make_shape(0, 0, 100, 50)
        s2 = _make_shape(200, 0, 200, 50)
        sr = ShapeRange([s1, s2])
        # bbox: left=0, right=400, center=200
        sr.align_center()
        assert s1.left == 150  # 200 - 100/2
        assert s2.left == 100  # 200 - 200/2

    def it_aligns_right(self):
        s1 = _make_shape(0, 0, 100, 50)
        s2 = _make_shape(200, 0, 200, 50)
        sr = ShapeRange([s1, s2])
        # bbox_right = 400
        sr.align_right()
        assert s1.left == 300  # 400 - 100
        assert s2.left == 200  # 400 - 200

    def it_aligns_top(self):
        s1 = _make_shape(0, 100, 50, 50)
        s2 = _make_shape(0, 300, 50, 80)
        sr = ShapeRange([s1, s2])
        sr.align_top()
        assert s1.top == 100
        assert s2.top == 100

    def it_aligns_middle(self):
        s1 = _make_shape(0, 0, 50, 100)
        s2 = _make_shape(0, 200, 50, 200)
        sr = ShapeRange([s1, s2])
        # bbox: top=0, bottom=400, middle=200
        sr.align_middle()
        assert s1.top == 150  # 200 - 100/2
        assert s2.top == 100  # 200 - 200/2

    def it_aligns_bottom(self):
        s1 = _make_shape(0, 0, 50, 100)
        s2 = _make_shape(0, 200, 50, 200)
        sr = ShapeRange([s1, s2])
        # bbox_bottom = 400
        sr.align_bottom()
        assert s1.top == 300  # 400 - 100
        assert s2.top == 200  # 400 - 200


class DescribeShapeRange_Distribution:
    """Unit tests for ShapeRange distribution methods."""

    def it_rejects_fewer_than_3_shapes_for_distribute_horizontal(self):
        shapes = [_make_shape(0, 0, 50, 50), _make_shape(100, 0, 50, 50)]
        sr = ShapeRange(shapes)
        with pytest.raises(ValueError, match="at least 3"):
            sr.distribute_horizontal()

    def it_rejects_fewer_than_3_shapes_for_distribute_vertical(self):
        shapes = [_make_shape(0, 0, 50, 50), _make_shape(0, 100, 50, 50)]
        sr = ShapeRange(shapes)
        with pytest.raises(ValueError, match="at least 3"):
            sr.distribute_vertical()

    def it_distributes_horizontally(self):
        # 3 shapes each 100 wide, spanning 0 to 500 (first at 0, last at 400)
        s1 = _make_shape(0, 0, 100, 50)
        s2 = _make_shape(100, 0, 100, 50)  # will be repositioned
        s3 = _make_shape(400, 0, 100, 50)
        sr = ShapeRange([s1, s2, s3])
        sr.distribute_horizontal()
        # total extent = 500 (0 to 500), total widths = 300, total gap = 200
        # gap = 100 each
        assert s1.left == 0
        assert s2.left == 200  # 0 + 100 + 100
        assert s3.left == 400  # 200 + 100 + 100

    def it_distributes_vertically(self):
        # 3 shapes each 100 tall, spanning 0 to 500 (first at 0, last at 400)
        s1 = _make_shape(0, 0, 50, 100)
        s2 = _make_shape(0, 100, 50, 100)  # will be repositioned
        s3 = _make_shape(0, 400, 50, 100)
        sr = ShapeRange([s1, s2, s3])
        sr.distribute_vertical()
        assert s1.top == 0
        assert s2.top == 200
        assert s3.top == 400

    def it_handles_varying_shape_sizes_in_distribution(self):
        # 3 shapes: widths 100, 50, 100 spanning 0..500
        s1 = _make_shape(0, 0, 100, 50)
        s2 = _make_shape(150, 0, 50, 50)
        s3 = _make_shape(400, 0, 100, 50)
        sr = ShapeRange([s1, s2, s3])
        sr.distribute_horizontal()
        # total extent = 500, total widths = 250, total gap = 250, gap = 125
        assert s1.left == 0
        assert s2.left == 225  # 0 + 100 + 125
        assert s3.left == 400  # 225 + 50 + 125


class DescribeShapeRange_BatchSetters:
    """Unit tests for ShapeRange batch property setters."""

    def it_sets_left_on_all_shapes(self):
        shapes = [_make_shape(0, 0, 50, 50), _make_shape(100, 0, 50, 50)]
        sr = ShapeRange(shapes)
        sr.set_left(Inches(1))
        for s in shapes:
            assert s.left == Inches(1)

    def it_sets_top_on_all_shapes(self):
        shapes = [_make_shape(0, 0, 50, 50), _make_shape(0, 100, 50, 50)]
        sr = ShapeRange(shapes)
        sr.set_top(Inches(2))
        for s in shapes:
            assert s.top == Inches(2)

    def it_sets_width_on_all_shapes(self):
        shapes = [_make_shape(0, 0, 50, 50), _make_shape(100, 0, 80, 50)]
        sr = ShapeRange(shapes)
        sr.set_width(Inches(3))
        for s in shapes:
            assert s.width == Inches(3)

    def it_sets_height_on_all_shapes(self):
        shapes = [_make_shape(0, 0, 50, 50), _make_shape(0, 100, 50, 80)]
        sr = ShapeRange(shapes)
        sr.set_height(Inches(1))
        for s in shapes:
            assert s.height == Inches(1)

    def it_sets_rotation_on_all_shapes(self):
        shapes = [_make_shape(0, 0, 50, 50), _make_shape(100, 0, 50, 50)]
        sr = ShapeRange(shapes)
        sr.set_rotation(45.0)
        for s in shapes:
            assert s.rotation == 45.0

    def it_sets_hidden_on_all_shapes(self):
        shapes = [_make_shape(0, 0, 50, 50), _make_shape(100, 0, 50, 50)]
        sr = ShapeRange(shapes)
        sr.set_hidden(True)
        for s in shapes:
            assert s.hidden is True
        sr.set_hidden(False)
        for s in shapes:
            assert s.hidden is False
