"""Unit-test suite for pptx.shapes.graphfrm module."""

from __future__ import annotations

import pytest

from pptx.chart.chart import Chart
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import ShapeError, UnsupportedEffectError
from pptx.parts.chart import ChartPart
from pptx.parts.chartex import ChartExPart
from pptx.parts.diagram import DiagramDataPart
from pptx.parts.embeddedpackage import EmbeddedPackagePart
from pptx.parts.slide import SlidePart
from pptx.shapes.graphfrm import GraphicFrame, _OleFormat
from pptx.shapes.shapetree import SlideShapes
from pptx.smartart import SmartArt
from pptx.spec import (
    GRAPHIC_DATA_URI_CHART,
    GRAPHIC_DATA_URI_CHARTEX,
    GRAPHIC_DATA_URI_DIAGRAM,
    GRAPHIC_DATA_URI_OLEOBJ,
    GRAPHIC_DATA_URI_TABLE,
)

from ..unitutil.cxml import element
from ..unitutil.mock import class_mock, instance_mock, property_mock


class DescribeGraphicFrame:
    """Unit-test suite for `pptx.shapes.graphfrm.GraphicFrame` object."""

    def it_provides_access_to_the_chart_it_contains(
        self, request, has_chart_prop_, chart_part_, chart_
    ):
        has_chart_prop_.return_value = True
        property_mock(request, GraphicFrame, "chart_part", return_value=chart_part_)
        chart_part_.chart = chart_

        assert GraphicFrame(None, None).chart is chart_

    def but_it_raises_on_chart_if_there_isnt_one(self, has_chart_prop_):
        has_chart_prop_.return_value = False

        with pytest.raises(ValueError) as e:
            GraphicFrame(None, None).chart
        assert str(e.value) == "shape does not contain a chart"

    def it_provides_access_to_its_chart_part(self, request, chart_part_):
        slide_part_ = instance_mock(request, SlidePart)
        slide_part_.related_part.return_value = chart_part_
        property_mock(request, GraphicFrame, "part", return_value=slide_part_)
        graphic_frame = GraphicFrame(
            element("p:graphicFrame/a:graphic/a:graphicData/c:chart{r:id=rId42}"), None
        )

        chart_part = graphic_frame.chart_part

        slide_part_.related_part.assert_called_once_with("rId42")
        assert chart_part is chart_part_

    @pytest.mark.parametrize(
        "graphicData_uri, expected_value",
        (
            (GRAPHIC_DATA_URI_CHART, True),
            (GRAPHIC_DATA_URI_OLEOBJ, False),
            (GRAPHIC_DATA_URI_TABLE, False),
        ),
    )
    def it_knows_whether_it_contains_a_chart(self, graphicData_uri, expected_value):
        graphicFrame = element("p:graphicFrame/a:graphic/a:graphicData{uri=%s}" % graphicData_uri)
        assert GraphicFrame(graphicFrame, None).has_chart is expected_value

    @pytest.mark.parametrize(
        "graphicData_uri, expected_value",
        (
            (GRAPHIC_DATA_URI_CHART, False),
            (GRAPHIC_DATA_URI_OLEOBJ, False),
            (GRAPHIC_DATA_URI_TABLE, True),
        ),
    )
    def it_knows_whether_it_contains_a_table(self, graphicData_uri, expected_value):
        graphicFrame = element("p:graphicFrame/a:graphic/a:graphicData{uri=%s}" % graphicData_uri)
        assert GraphicFrame(graphicFrame, None).has_table is expected_value

    @pytest.mark.parametrize(
        "graphicData_uri, expected_value",
        (
            (GRAPHIC_DATA_URI_CHART, False),
            (GRAPHIC_DATA_URI_CHARTEX, True),
            (GRAPHIC_DATA_URI_DIAGRAM, False),
        ),
    )
    def it_knows_whether_it_contains_a_chartex_chart(self, graphicData_uri, expected_value):
        graphicFrame = element("p:graphicFrame/a:graphic/a:graphicData{uri=%s}" % graphicData_uri)
        graphic_frame = GraphicFrame(graphicFrame, None)

        assert graphic_frame.has_chartex is expected_value
        if graphicData_uri == GRAPHIC_DATA_URI_CHARTEX:
            assert graphic_frame.has_chart is False

    def it_provides_access_to_the_chartex_chart_it_contains(self, request):
        chartex_part_ = instance_mock(request, ChartExPart)
        slide_part_ = instance_mock(request, SlidePart)
        slide_part_.related_part.return_value = chartex_part_
        property_mock(request, GraphicFrame, "part", return_value=slide_part_)
        graphic_frame = GraphicFrame(
            element(
                "p:graphicFrame/a:graphic/a:graphicData{uri=%s}/cx:chart{r:id=rId5}"
                % GRAPHIC_DATA_URI_CHARTEX
            ),
            None,
        )

        chartex = graphic_frame.chartex

        slide_part_.related_part.assert_called_once_with("rId5")
        assert chartex is chartex_part_.chartex

    def but_it_raises_on_chartex_when_it_does_not_contain_one(self):
        graphic_frame = GraphicFrame(
            element("p:graphicFrame/a:graphic/a:graphicData{uri=%s}" % GRAPHIC_DATA_URI_CHART),
            None,
        )
        with pytest.raises(ValueError) as e:
            graphic_frame.chartex
        assert str(e.value) == "shape does not contain a chartex chart"

    @pytest.mark.parametrize(
        "graphicData_uri, expected_value",
        (
            (GRAPHIC_DATA_URI_CHART, False),
            (GRAPHIC_DATA_URI_DIAGRAM, True),
            (GRAPHIC_DATA_URI_OLEOBJ, False),
            (GRAPHIC_DATA_URI_TABLE, False),
        ),
    )
    def it_knows_whether_it_contains_smartart(self, graphicData_uri, expected_value):
        graphicFrame = element("p:graphicFrame/a:graphic/a:graphicData{uri=%s}" % graphicData_uri)
        assert GraphicFrame(graphicFrame, None).has_smartart is expected_value

    def it_provides_access_to_the_smartart_it_contains(self, request):
        data_part_ = instance_mock(request, DiagramDataPart)
        smartart_ = instance_mock(request, SmartArt)
        SmartArt_ = class_mock(request, "pptx.shapes.graphfrm.SmartArt", return_value=smartart_)
        slide_part_ = instance_mock(request, SlidePart)
        slide_part_.related_part.return_value = data_part_
        property_mock(request, GraphicFrame, "part", return_value=slide_part_)
        graphic_frame = GraphicFrame(
            element(
                "p:graphicFrame/a:graphic/a:graphicData{uri=%s}"
                "/dgm:relIds{r:dm=rId7,r:lo=rId8,r:qs=rId9,r:cs=rId10}" % GRAPHIC_DATA_URI_DIAGRAM
            ),
            None,
        )

        smartart = graphic_frame.smartart

        slide_part_.related_part.assert_called_once_with("rId7")
        SmartArt_.assert_called_once_with(data_part_, slide_part_)
        assert smartart is smartart_

    def but_it_raises_on_smartart_when_it_does_not_contain_smartart(self):
        graphic_frame = GraphicFrame(
            element("p:graphicFrame/a:graphic/a:graphicData{uri=%s}" % GRAPHIC_DATA_URI_TABLE),
            None,
        )
        with pytest.raises(ValueError) as e:
            graphic_frame.smartart
        assert str(e.value) == "shape does not contain SmartArt"

    def it_provides_access_to_the_OleFormat_object(self, request):
        ole_format_ = instance_mock(request, _OleFormat)
        _OleFormat_ = class_mock(
            request, "pptx.shapes.graphfrm._OleFormat", return_value=ole_format_
        )
        graphicFrame = element(
            "p:graphicFrame/a:graphic/a:graphicData{uri=http://schemas.openxmlformats"
            ".org/presentationml/2006/ole}"
        )
        parent_ = instance_mock(request, SlideShapes)
        graphic_frame = GraphicFrame(graphicFrame, parent_)

        ole_format = graphic_frame.ole_format

        _OleFormat_.assert_called_once_with(graphicFrame.graphicData, parent_)
        assert ole_format is ole_format_

    def but_it_raises_on_ole_format_when_this_is_not_an_OLE_object(self):
        graphic_frame = GraphicFrame(
            element(
                "p:graphicFrame/a:graphic/a:graphicData{uri=http://schemas.openxmlfor"
                "mats.org/drawingml/2006/table}"
            ),
            None,
        )
        with pytest.raises(ValueError) as e:
            graphic_frame.ole_format
        assert str(e.value) == "not an OLE-object shape"

    def it_reads_a_table_frames_effects_without_changing_the_xml(self):
        frame = _slide_with_table_and_chart().shapes[0]
        tbl = frame.table._tbl
        tbl.remove(tbl.tblPr)

        assert frame.shadow.inherit is True
        assert frame.glow.radius is None
        assert tbl.tblPr is None

    def and_it_writes_them_to_the_tables_tblPr(self):
        frame = _slide_with_table_and_chart().shapes[0]

        frame.shadow.inherit = False

        tblPr = frame.table._tbl.tblPr
        assert tblPr.effectLst is not None
        assert [child.tag.split("}")[1] for child in tblPr] == ["effectLst", "tableStyleId"]

    def it_reads_a_chart_frames_effects_without_changing_the_chart_part(self):
        frame = _slide_with_table_and_chart().shapes[1]
        chartSpace = frame.chart._chartSpace

        assert frame.shadow.inherit is True
        assert frame.soft_edge.radius is None
        assert frame.three_d.extrusion_height is None
        assert chartSpace.spPr is None

    def and_it_writes_them_and_3D_to_the_chart_spaces_spPr(self):
        frame = _slide_with_table_and_chart().shapes[1]

        frame.shadow.inherit = False
        frame.three_d.extrusion_height = 12700

        spPr = frame.chart._chartSpace.spPr
        assert spPr.effectLst is not None
        assert spPr.sp3d is not None

    def but_it_has_no_3D_formatting_for_a_table(self):
        frame = _slide_with_table_and_chart().shapes[0]
        with pytest.raises(UnsupportedEffectError, match="only supported for a graphic frame containing a chart"):
            frame.three_d

    @pytest.mark.parametrize("name", ["shadow", "glow", "reflection", "soft_edge", "three_d"])
    def but_it_raises_UnsupportedEffectError_for_other_content(self, name: str):
        graphicFrame = element(
            "p:graphicFrame/a:graphic/a:graphicData{uri=http://schemas.openxmlformats.org/"
            "drawingml/2006/diagram}"
        )
        frame = GraphicFrame(graphicFrame, None)

        with pytest.raises(UnsupportedEffectError, match="only supported for a graphic frame"):
            getattr(frame, name)
        # -- still what these raised before: NotImplementedError (shadow), AttributeError --
        with pytest.raises(NotImplementedError):
            getattr(frame, name)
        assert hasattr(frame, name) is False

    def it_round_trips_a_table_shadow_and_a_chart_shadow(self):
        import io

        from pptx import Presentation

        slide = _slide_with_table_and_chart()
        for frame in slide.shapes:
            frame.shadow.inherit = False
        stream = io.BytesIO()
        slide.part.package.save(stream)
        stream.seek(0)

        frames = Presentation(stream).slides[0].shapes

        assert [frame.shadow.inherit for frame in frames] == [False, False]
        tblPr = frames[0].table._tbl.tblPr
        assert [child.tag.split("}")[1] for child in tblPr] == ["effectLst", "tableStyleId"]

    @pytest.mark.parametrize(
        "uri, oleObj_child, expected_value",
        (
            (GRAPHIC_DATA_URI_CHART, None, MSO_SHAPE_TYPE.CHART),
            (GRAPHIC_DATA_URI_OLEOBJ, "embed", MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT),
            (GRAPHIC_DATA_URI_OLEOBJ, "link", MSO_SHAPE_TYPE.LINKED_OLE_OBJECT),
            (GRAPHIC_DATA_URI_TABLE, None, MSO_SHAPE_TYPE.TABLE),
            (GRAPHIC_DATA_URI_DIAGRAM, None, MSO_SHAPE_TYPE.SMART_ART),
            (GRAPHIC_DATA_URI_CHARTEX, None, MSO_SHAPE_TYPE.CHART),
            ("foobar", None, None),
        ),
    )
    def it_knows_its_shape_type(self, uri, oleObj_child, expected_value):
        graphicFrame = element(
            ("p:graphicFrame/a:graphic/a:graphicData{uri=%s}/p:oleObj/p:%s" % (uri, oleObj_child))
            if oleObj_child
            else "p:graphicFrame/a:graphic/a:graphicData{uri=%s}" % uri
        )
        assert GraphicFrame(graphicFrame, None).shape_type is expected_value

    # fixture components ---------------------------------------------

    @pytest.fixture
    def chart_(self, request):
        return instance_mock(request, Chart)

    @pytest.fixture
    def chart_part_(self, request, chart_):
        return instance_mock(request, ChartPart, chart=chart_)

    @pytest.fixture
    def has_chart_prop_(self, request):
        return property_mock(request, GraphicFrame, "has_chart")


class Describe_OleFormat:
    """Unit-test suite for `pptx.shapes.graphfrm._OleFormat` object."""

    def it_provides_access_to_the_OLE_object_blob(self, request):
        ole_obj_part_ = instance_mock(request, EmbeddedPackagePart, blob=b"0123456789")
        slide_part_ = instance_mock(request, SlidePart)
        slide_part_.related_part.return_value = ole_obj_part_
        property_mock(request, _OleFormat, "part", return_value=slide_part_)
        ole_format = _OleFormat(element("a:graphicData/p:oleObj{r:id=rId7}"), None)

        blob = ole_format.blob

        slide_part_.related_part.assert_called_once_with("rId7")
        assert blob == b"0123456789"

    def it_knows_the_OLE_object_prog_id(self):
        graphicData = element("a:graphicData/p:oleObj{progId=Excel.Sheet.12}")
        assert _OleFormat(graphicData, None).prog_id == "Excel.Sheet.12"

    def it_knows_whether_to_show_the_OLE_object_as_an_icon(self):
        graphicData = element("a:graphicData/p:oleObj{showAsIcon=1}")
        assert _OleFormat(graphicData, None).show_as_icon is True


def _slide_with_table_and_chart():
    """A slide holding a table frame then a chart frame, built through the public API."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_table(2, 2, 0, 0, Inches(3), Inches(1))
    chart_data = CategoryChartData()
    chart_data.categories = ["a", "b"]
    chart_data.add_series("s", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, 0, Inches(2), Inches(4), Inches(3), chart_data
    )
    return slide
