# pyright: reportPrivateUsage=false

"""Round-trip integration tests for python-pptx.

These tests create presentations, save to BytesIO, reload, and verify
that content survives the round-trip through serialization/deserialization.
"""

from __future__ import annotations

import subprocess
import sys
from io import BytesIO
from textwrap import dedent

import pytest


def _run_roundtrip_test(test_code: str) -> None:
    """Run a round-trip test in a subprocess to avoid mock pollution from other tests."""
    preamble = "from io import BytesIO\nfrom pptx import Presentation\nfrom pptx.util import Inches, Pt\n"
    script = preamble + dedent(test_code)
    result = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True)
    if result.returncode != 0:
        pytest.fail(f"Round-trip test failed:\\n{result.stderr}")


class DescribeRoundTrip:
    """Integration tests verifying save/load round-trip fidelity."""

    def it_round_trips_a_blank_presentation(self):
        _run_roundtrip_test("""\
prs = Presentation()
stream = BytesIO()
prs.save(stream)
stream.seek(0)
prs2 = Presentation(stream)
assert len(prs2.slides) == 0
assert len(prs2.slide_layouts) > 0
""")

    def it_round_trips_a_slide_with_shapes(self):
        _run_roundtrip_test("""\
prs = Presentation()
layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(layout)
txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
txBox.text_frame.text = "Hello, World!"

stream = BytesIO()
prs.save(stream)
stream.seek(0)
prs2 = Presentation(stream)

assert len(prs2.slides) == 1
slide2 = prs2.slides[0]
textboxes = [s for s in slide2.shapes if s.has_text_frame and not s.is_placeholder]
assert len(textboxes) == 1
assert textboxes[0].text_frame.text == "Hello, World!"
""")

    def it_round_trips_shape_accessibility_properties(self):
        _run_roundtrip_test("""\
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])
shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
shape.alternative_text = "A test textbox"
shape.title = "Test Shape"
shape.decorative = True

stream = BytesIO()
prs.save(stream)
stream.seek(0)
prs2 = Presentation(stream)

shape2 = [s for s in prs2.slides[0].shapes if not s.is_placeholder][0]
assert shape2.alternative_text == "A test textbox"
assert shape2.title == "Test Shape"
assert shape2.decorative is True
""")

    def it_round_trips_font_properties(self):
        _run_roundtrip_test("""\
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])
shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
tf = shape.text_frame
tf.text = "Styled text"
run = tf.paragraphs[0].runs[0]
run.font.bold = True
run.font.italic = True
run.font.size = Pt(24)
run.font.name = "Arial"

stream = BytesIO()
prs.save(stream)
stream.seek(0)
prs2 = Presentation(stream)

shape2 = [s for s in prs2.slides[0].shapes if not s.is_placeholder][0]
run2 = shape2.text_frame.paragraphs[0].runs[0]
assert run2.font.bold is True
assert run2.font.italic is True
assert run2.font.size == Pt(24)
assert run2.font.name == "Arial"
""")

    def it_round_trips_slide_operations(self):
        _run_roundtrip_test("""\
prs = Presentation()
layout = prs.slide_layouts[5]

slide1 = prs.slides.add_slide(layout)
slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "Slide 1"
slide2 = prs.slides.add_slide(layout)
slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "Slide 2"
slide3 = prs.slides.add_slide(layout)
slide3.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "Slide 3"

prs.slides.delete(slide2)
assert len(prs.slides) == 2

stream = BytesIO()
prs.save(stream)
stream.seek(0)
prs2 = Presentation(stream)

assert len(prs2.slides) == 2
texts = []
for s in prs2.slides:
    for shape in s.shapes:
        if shape.has_text_frame and not shape.is_placeholder:
            texts.append(shape.text_frame.text)
assert texts == ["Slide 1", "Slide 3"]
""")

    def it_round_trips_a_duplicated_slide(self):
        _run_roundtrip_test("""\
prs = Presentation()
layout = prs.slide_layouts[5]
slide = prs.slides.add_slide(layout)
shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
shape.text_frame.text = "Original"

new_slide = prs.slides.duplicate(slide)
assert len(prs.slides) == 2

stream = BytesIO()
prs.save(stream)
stream.seek(0)
prs2 = Presentation(stream)

assert len(prs2.slides) == 2
for s in prs2.slides:
    textboxes = [sh for sh in s.shapes if sh.has_text_frame and not sh.is_placeholder]
    assert len(textboxes) == 1
    assert textboxes[0].text_frame.text == "Original"
""")

    def it_gives_a_duplicated_slide_its_own_chart(self):
        _run_roundtrip_test("""\
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
chart_data = CategoryChartData()
chart_data.categories = ["a", "b"]
chart_data.add_series("S1", (1, 2))
slide.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), chart_data
)
dup = prs.slides.duplicate(slide)
assert slide.shapes[0].chart.part is not dup.shapes[0].chart.part

new_data = CategoryChartData()
new_data.categories = ["a", "b"]
new_data.add_series("S1", (7, 8))
dup.shapes[0].chart.replace_data(new_data)

stream = BytesIO()
prs.save(stream)
stream.seek(0)
prs2 = Presentation(stream)

values = [tuple(s.shapes[0].chart.plots[0].series[0].values) for s in prs2.slides]
assert values == [(1.0, 2.0), (7.0, 8.0)], values
partnames = {s.shapes[0].chart.part.partname for s in prs2.slides}
assert len(partnames) == 2, partnames
""")

    def it_keeps_the_workbook_of_an_imported_chart(self):
        _run_roundtrip_test("""\
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

src = Presentation()
src_slide = src.slides.add_slide(src.slide_layouts[6])
for top in (Inches(0), Inches(3)):
    chart_data = CategoryChartData()
    chart_data.categories = ["a", "b"]
    chart_data.add_series("S1", (1, 2) if top == 0 else (3, 4))
    src_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), top, Inches(4), Inches(2), chart_data
    )

dst = Presentation()
dst.slides.import_slide(src.slides[0])

stream = BytesIO()
dst.save(stream)
stream.seek(0)
prs2 = Presentation(stream)

charts = [shape.chart for shape in prs2.slides[0].shapes]
assert [tuple(c.plots[0].series[0].values) for c in charts] == [(1.0, 2.0), (3.0, 4.0)]
workbooks = [c.part.chart_workbook.xlsx_part for c in charts]
assert all(wb is not None for wb in workbooks)
assert len({wb.partname for wb in workbooks}) == 2
assert len({c.part.partname for c in charts}) == 2
""")

    def it_round_trips_hidden_shapes(self):
        _run_roundtrip_test("""\
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])
shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
shape.hidden = True

stream = BytesIO()
prs.save(stream)
stream.seek(0)
prs2 = Presentation(stream)

shape2 = [s for s in prs2.slides[0].shapes if not s.is_placeholder][0]
assert shape2.hidden is True
""")

    def it_round_trips_a_duplicated_shape(self):
        _run_roundtrip_test("""\
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])
shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
shape.text_frame.text = "Original"

new_shape = slide.shapes.duplicate_shape(shape)
assert len([s for s in slide.shapes if not s.is_placeholder]) == 2
assert new_shape.text_frame.text == "Original"
assert new_shape.shape_id != shape.shape_id

stream = BytesIO()
prs.save(stream)
stream.seek(0)
prs2 = Presentation(stream)

slide2 = prs2.slides[0]
textboxes = [s for s in slide2.shapes if s.has_text_frame and not s.is_placeholder]
assert len(textboxes) == 2
assert all(t.text_frame.text == "Original" for t in textboxes)
# shape IDs should be unique
ids = [t.shape_id for t in textboxes]
assert len(set(ids)) == 2
""")

    def it_round_trips_slide_numbers(self):
        _run_roundtrip_test("""\
prs = Presentation()
prs.first_slide_number = 5

stream = BytesIO()
prs.save(stream)
stream.seek(0)
prs2 = Presentation(stream)

assert prs2.first_slide_number == 5
""")

    def it_round_trips_transparency_color(self):
        _run_roundtrip_test("""\
from pptx.dml.color import RGBColor

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

import os, tempfile
# create a tiny valid PNG (1x1 white pixel)
png_bytes = (
    b'\\x89PNG\\r\\n\\x1a\\n\\x00\\x00\\x00\\rIHDR\\x00\\x00\\x00\\x01'
    b'\\x00\\x00\\x00\\x01\\x08\\x02\\x00\\x00\\x00\\x90wS\\xde\\x00'
    b'\\x00\\x00\\x0cIDATx\\x9cc\\xf8\\x0f\\x00\\x00\\x01\\x01\\x00'
    b'\\x05\\x18\\xd8N\\x00\\x00\\x00\\x00IEND\\xaeB`\\x82'
)
with tempfile.NamedTemporaryFile(suffix='.png', delete=False) as f:
    f.write(png_bytes)
    img_path = f.name

try:
    pic = slide.shapes.add_picture(img_path, Inches(1), Inches(1))
    pic.transparency_color = RGBColor(0xFF, 0xFF, 0xFF)
    assert pic.transparency_color == RGBColor(0xFF, 0xFF, 0xFF)

    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    prs2 = Presentation(stream)

    slide2 = prs2.slides[0]
    pics = [s for s in slide2.shapes if s.shape_type is not None]
    from pptx.shapes.picture import Picture
    pic2 = [s for s in slide2.shapes if isinstance(s, Picture)][0]
    assert pic2.transparency_color == RGBColor(0xFF, 0xFF, 0xFF)

    # can remove it
    pic2.transparency_color = None
    assert pic2.transparency_color is None
finally:
    os.unlink(img_path)
""")

    def it_round_trips_theme_effect_scheme(self):
        _run_roundtrip_test("""\
from pptx.theme import EffectScheme

prs = Presentation()

stream = BytesIO()
prs.save(stream)
stream.seek(0)
prs2 = Presentation(stream)

master = prs2.slide_masters[0]
theme = master.theme
effect_scheme = theme.effect_scheme
assert effect_scheme is not None
assert len(effect_scheme) == 3
assert effect_scheme.subtle.has_effect_list is True
assert effect_scheme.intense.has_3d_scene is True
assert effect_scheme.intense.has_3d_shape is True
assert effect_scheme.name == "Office"
""")

    def it_round_trips_slide_import(self):
        _run_roundtrip_test("""\
# --- create source presentation with a text box ---
src_prs = Presentation()
src_slide = src_prs.slides.add_slide(src_prs.slide_layouts[6])  # blank layout
txBox = src_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
txBox.text_frame.text = "Imported Slide"

# --- save and reload source ---
src_stream = BytesIO()
src_prs.save(src_stream)
src_stream.seek(0)
src_prs2 = Presentation(src_stream)

# --- create target presentation and import slide ---
tgt_prs = Presentation()
imported = tgt_prs.slides.import_slide(src_prs2.slides[0])

assert imported is not None
# target should now have 1 slide (the imported one)
assert len(tgt_prs.slides) == 1

# verify content survived
shapes = list(tgt_prs.slides[0].shapes)
textboxes = [s for s in shapes if s.has_text_frame and not s.is_placeholder]
assert len(textboxes) == 1
assert textboxes[0].text_frame.text == "Imported Slide"

# --- save and reload to verify round-trip ---
tgt_stream = BytesIO()
tgt_prs.save(tgt_stream)
tgt_stream.seek(0)
tgt_prs2 = Presentation(tgt_stream)

assert len(tgt_prs2.slides) == 1
shapes2 = list(tgt_prs2.slides[0].shapes)
textboxes2 = [s for s in shapes2 if s.has_text_frame and not s.is_placeholder]
assert len(textboxes2) == 1
assert textboxes2[0].text_frame.text == "Imported Slide"
""")

    def it_round_trips_merge_presentations(self):
        _run_roundtrip_test("""\
# --- create two source presentations ---
prs_a = Presentation()
slide_a = prs_a.slides.add_slide(prs_a.slide_layouts[6])
txBox_a = slide_a.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
txBox_a.text_frame.text = "Slide A"

prs_b = Presentation()
slide_b1 = prs_b.slides.add_slide(prs_b.slide_layouts[6])
txBox_b1 = slide_b1.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
txBox_b1.text_frame.text = "Slide B1"
slide_b2 = prs_b.slides.add_slide(prs_b.slide_layouts[6])
txBox_b2 = slide_b2.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
txBox_b2.text_frame.text = "Slide B2"

# --- save and reload ---
stream_a = BytesIO()
prs_a.save(stream_a)
stream_a.seek(0)
prs_a2 = Presentation(stream_a)

stream_b = BytesIO()
prs_b.save(stream_b)
stream_b.seek(0)
prs_b2 = Presentation(stream_b)

# --- merge B into A ---
new_slides = prs_a2.slides.merge(prs_b2)
assert len(new_slides) == 2
assert len(prs_a2.slides) == 3  # 1 original + 2 merged

# --- save and reload ---
out_stream = BytesIO()
prs_a2.save(out_stream)
out_stream.seek(0)
merged = Presentation(out_stream)

assert len(merged.slides) == 3
texts = []
for slide in merged.slides:
    for shape in slide.shapes:
        if shape.has_text_frame and not shape.is_placeholder:
            texts.append(shape.text_frame.text)
assert texts == ["Slide A", "Slide B1", "Slide B2"]
""")

    def it_preserves_potx_content_type_on_round_trip(self):
        _run_roundtrip_test("""\
from pptx.opc.constants import CONTENT_TYPE as CT

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Template Slide"

# Override the main part content type to template
del prs.part.__dict__['content_type']  # clear lazyproperty cache
prs.part._content_type = CT.PML_TEMPLATE_MAIN

stream = BytesIO()
prs.save(stream)
stream.seek(0)

# Reload and verify content type is preserved
prs2 = Presentation(stream)
assert prs2.part.content_type == CT.PML_TEMPLATE_MAIN
assert prs2.slides[0].shapes.title.text == "Template Slide"
""")

    def it_round_trips_shape_range_alignment(self):
        _run_roundtrip_test("""\
from pptx.shapes.range import ShapeRange

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])

s1 = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
s2 = slide.shapes.add_textbox(Inches(3), Inches(2), Inches(2), Inches(1))
s3 = slide.shapes.add_textbox(Inches(5), Inches(3), Inches(2), Inches(1))

# use shape_range convenience method
sr = slide.shapes.shape_range()
assert len(sr) == 3

# align left and distribute vertical
sr.align_left()
sr.distribute_vertical()

stream = BytesIO()
prs.save(stream)
stream.seek(0)
prs2 = Presentation(stream)

slide2 = prs2.slides[0]
shapes = [s for s in slide2.shapes if not s.is_placeholder]
assert len(shapes) == 3

# all shapes should have the same left edge after align_left
lefts = [s.left for s in shapes]
assert len(set(lefts)) == 1
""")

    def it_imports_slide_with_image(self):
        _run_roundtrip_test("""\
import os, tempfile
from pptx.shapes.picture import Picture

# tiny valid PNG (1x1 white pixel)
png_bytes = (
    b'\\x89PNG\\r\\n\\x1a\\n\\x00\\x00\\x00\\rIHDR\\x00\\x00\\x00\\x01'
    b'\\x00\\x00\\x00\\x01\\x08\\x02\\x00\\x00\\x00\\x90wS\\xde\\x00'
    b'\\x00\\x00\\x0cIDATx\\x9cc\\xf8\\x0f\\x00\\x00\\x01\\x01\\x00'
    b'\\x05\\x18\\xd8N\\x00\\x00\\x00\\x00IEND\\xaeB`\\x82'
)
with tempfile.NamedTemporaryFile(suffix='.png', delete=False) as f:
    f.write(png_bytes)
    img_path = f.name

try:
    # --- source with image ---
    src = Presentation()
    slide = src.slides.add_slide(src.slide_layouts[6])
    slide.shapes.add_picture(img_path, Inches(1), Inches(1))

    src_stream = BytesIO()
    src.save(src_stream)
    src_stream.seek(0)
    src2 = Presentation(src_stream)

    # --- import into target ---
    tgt = Presentation()
    tgt.slides.import_slide(src2.slides[0])

    tgt_stream = BytesIO()
    tgt.save(tgt_stream)
    tgt_stream.seek(0)
    tgt2 = Presentation(tgt_stream)

    assert len(tgt2.slides) == 1
    pics = [s for s in tgt2.slides[0].shapes if isinstance(s, Picture)]
    assert len(pics) == 1
finally:
    os.unlink(img_path)
""")

