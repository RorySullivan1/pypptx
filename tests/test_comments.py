# pyright: reportPrivateUsage=false

"""Unit-test suite for slide comments API."""

from __future__ import annotations

import datetime as dt

import pytest

from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.parts.comments import (
    AuthorsPart,
    CommentAuthorsPart,
    CommentsPart,
    ModernCommentsPart,
)
from pptx.parts.presentation import PresentationPart
from pptx.slide import Comment, CommentReply, SlideComments, ThreadedComment

from .unitutil.cxml import element
from .unitutil.mock import instance_mock



class DescribeCommentAuthorsPart:
    """Unit-test suite for `pptx.parts.comments.CommentAuthorsPart`."""

    def _part(self):
        xml = '<p:cmAuthorLst %s/>' % nsdecls("p")
        element = parse_xml(xml)
        return CommentAuthorsPart("/ppt/commentAuthors.xml", "application/test", None, element)

    def it_starts_empty(self):
        part = self._part()
        assert len(part) == 0
        assert list(part) == []

    def it_can_add_an_author(self):
        part = self._part()
        author = part.get_or_add_author("Alice", "A")
        assert author.name == "Alice"
        assert author.initials == "A"
        assert author.id == 0
        assert len(part) == 1

    def it_reuses_existing_authors(self):
        part = self._part()
        a1 = part.get_or_add_author("Alice", "A")
        a2 = part.get_or_add_author("Alice", "A")
        assert a1 is a2
        assert len(part) == 1

    def it_assigns_unique_ids(self):
        part = self._part()
        a1 = part.get_or_add_author("Alice", "A")
        a2 = part.get_or_add_author("Bob", "B")
        assert a1.id == 0
        assert a2.id == 1

    def it_can_find_an_author_by_id(self):
        part = self._part()
        part.get_or_add_author("Alice", "A")
        author = part.get_author(0)
        assert author is not None
        assert author.name == "Alice"
        assert part.get_author(99) is None


class DescribeCommentsPart:
    """Unit-test suite for `pptx.parts.comments.CommentsPart`."""

    def _part(self, n_comments=0):
        parts = ['<p:cmLst %s>' % nsdecls("p")]
        for i in range(n_comments):
            parts.append(
                f'<p:cm authorId="0" idx="{i}">'
                f'<p:pos x="100" y="200"/>'
                f'<p:text>Comment {i}</p:text>'
                f'</p:cm>'
            )
        parts.append('</p:cmLst>')
        element = parse_xml("".join(parts))
        return CommentsPart("/ppt/comments/comment1.xml", "application/test", None, element)

    def it_starts_empty(self):
        part = self._part()
        assert len(part) == 0

    def it_counts_comments(self):
        part = self._part(n_comments=3)
        assert len(part) == 3

    def it_iterates_comments(self):
        part = self._part(n_comments=2)
        comments = list(part)
        assert len(comments) == 2


class DescribeAuthorsPart:
    """Unit-test suite for `pptx.parts.comments.AuthorsPart`."""

    def it_provides_access_to_its_authors(self):
        part = _authors_part()
        assert len(part) == 2
        assert [a.name for a in part] == ["Alice", "Bob"]
        assert [a.name for a in part.authors] == ["Alice", "Bob"]

    def it_reuses_an_author_with_the_same_name(self):
        part = _authors_part()

        assert part.get_or_add_author("Bob", "XX").id == "B2"
        assert len(part) == 2

    def it_adds_an_author_it_does_not_have(self):
        part = _authors_part()

        carol = part.get_or_add_author("Carol", "CC")

        assert (carol.name, carol.initials) == ("Carol", "CC")
        assert len(part) == 3

    def it_can_find_an_author_by_id(self):
        part = _authors_part()
        author = part.get_author("B2")
        assert author is not None
        assert author.name == "Bob"
        assert part.get_author("ZZ") is None


class DescribeModernCommentsPart:
    """Unit-test suite for `pptx.parts.comments.ModernCommentsPart`."""

    def it_provides_access_to_its_threads(self):
        cmLst = element(
            "p188:cmLst/(p188:cm{id=C1,authorId=A1,created=2024-05-01T09:30:00.000}"
            ",p188:cm{id=C2,authorId=A1,created=2024-05-02T09:30:00.000})"
        )
        part = ModernCommentsPart(
            "/ppt/comments/modernComment_100_1.xml", "application/test", None, cmLst
        )

        assert len(part) == 2
        assert [cm.id for cm in part] == ["C1", "C2"]
        assert [cm.id for cm in part.comments] == ["C1", "C2"]


class DescribeComment:
    """Unit-test suite for `pptx.slide.Comment`."""

    def _make_comment(self, author_name="Alice", text="Hello", x=100, y=200):
        # Build authors part
        authors_xml = (
            '<p:cmAuthorLst %s>'
            '<p:cmAuthor id="0" name="%s" initials="A" lastIdx="0" clrIdx="0"/>'
            '</p:cmAuthorLst>'
        ) % (nsdecls("p"), author_name)
        authors_element = parse_xml(authors_xml)
        authors_part = CommentAuthorsPart(
            "/ppt/commentAuthors.xml", "application/test", None, authors_element
        )

        # Build comment element
        cm_xml = (
            '<p:cm %s authorId="0" idx="0">'
            '<p:pos x="%d" y="%d"/>'
            '<p:text>%s</p:text>'
            '</p:cm>'
        ) % (nsdecls("p"), x, y, text)
        cm = parse_xml(cm_xml)
        return Comment(cm, authors_part)

    def it_knows_its_author(self):
        comment = self._make_comment(author_name="Alice")
        assert comment.author == "Alice"

    def it_knows_its_text(self):
        comment = self._make_comment(text="Hello world")
        assert comment.text == "Hello world"

    def it_can_change_its_text(self):
        comment = self._make_comment(text="Old")
        comment.text = "New"
        assert comment.text == "New"

    def it_knows_its_position(self):
        comment = self._make_comment(x=100, y=200)
        assert comment.position == (100, 200)

    def it_can_be_deleted(self):
        # Build a comment list with one comment
        xml = (
            '<p:cmLst %s>'
            '<p:cm authorId="0" idx="0">'
            '<p:pos x="0" y="0"/><p:text>Test</p:text>'
            '</p:cm>'
            '</p:cmLst>'
        ) % nsdecls("p")
        cmLst = parse_xml(xml)
        cm = cmLst[0]
        authors_xml = '<p:cmAuthorLst %s/>' % nsdecls("p")
        authors_part = CommentAuthorsPart(
            "/ppt/commentAuthors.xml", "application/test", None, parse_xml(authors_xml)
        )
        comment = Comment(cm, authors_part)
        comment.delete()
        assert len(cmLst) == 0


def _authors_part() -> AuthorsPart:
    authorLst = element(
        "p188:authorLst/(p188:author{id=A1,name=Alice,initials=AA,userId=alice,providerId=None}"
        ",p188:author{id=B2,name=Bob,userId=bob,providerId=AD})"
    )
    return AuthorsPart("/ppt/authors.xml", "application/test", None, authorLst)


def _thread_element(attrs: str = "", children: str = ""):
    """`p188:cm` authored by Alice (`A1`); `attrs` and `children` are cxml fragments."""
    return element(
        "p188:cm{id=C1,authorId=A1,created=2024-05-01T09:30:00.250%s}%s" % (attrs, children)
    )


def _txBody(text: str) -> str:
    return '/p188:txBody/(a:bodyPr,a:p/a:r/a:t"%s")' % text


class DescribeThreadedComment:
    """Unit-test suite for `pptx.slide.ThreadedComment`."""

    @pytest.fixture
    def prs_part_(self, request: pytest.FixtureRequest):
        return instance_mock(request, PresentationPart, authors_part=_authors_part())

    def it_knows_its_id_text_and_created_time(self, prs_part_):
        thread = ThreadedComment(_thread_element(children=_txBody("Fix")), prs_part_)

        assert thread.id == "C1"
        assert thread.text == "Fix"
        assert thread.created == dt.datetime(2024, 5, 1, 9, 30, 0, 250000)

    def it_knows_its_author(self, prs_part_):
        thread = ThreadedComment(_thread_element(), prs_part_)

        assert thread.author == "Alice"
        assert thread.author_initials == "AA"

    def but_it_has_a_blank_author_when_not_in_the_authors_part(self, prs_part_):
        cm = _thread_element()
        cm.set("authorId", "ZZ")
        assert ThreadedComment(cm, prs_part_).author == ""

        prs_part_.authors_part = None
        cm.set("authorId", "A1")
        assert ThreadedComment(cm, prs_part_).author == ""
        assert ThreadedComment(cm, prs_part_).author_initials == ""

    @pytest.mark.parametrize(
        ("attrs", "expected_status", "expected_resolved"),
        [
            ("", "active", False),
            (",status=resolved", "resolved", True),
            (",status=closed", "closed", False),
        ],
    )
    def it_knows_whether_it_is_resolved(
        self, attrs: str, expected_status: str, expected_resolved: bool, prs_part_
    ):
        thread = ThreadedComment(_thread_element(attrs), prs_part_)

        assert thread.status == expected_status
        assert thread.resolved is expected_resolved

    @pytest.mark.parametrize(
        ("children", "expected_value"),
        [
            ("/pc:sldMkLst/(pc:docMk,pc:sldMk{sldId=256})", "slide"),
            ("/ac:deMkLst/(pc:docMk,pc:sldMk{sldId=256})", "shape"),
            ("/ac:txMkLst", "text"),
            ("/p188:unknownAnchor", "unknown"),
        ],
    )
    def it_knows_what_it_is_anchored_to(self, prs_part_, children: str, expected_value: str):
        thread = ThreadedComment(_thread_element(children=children), prs_part_)
        assert thread.anchor == expected_value

    @pytest.mark.parametrize(
        ("children", "expected_value"),
        [("/p188:pos{x=100,y=200}", (100, 200)), ("", None)],
    )
    def it_knows_its_position(self, prs_part_, children: str, expected_value):
        thread = ThreadedComment(_thread_element(children=children), prs_part_)
        assert thread.position == expected_value

    def it_provides_access_to_its_replies_in_order(self, prs_part_):
        children = (
            "/p188:replyLst/("
            "p188:reply{id=R1,authorId=B2,created=2024-05-01T10:00:00Z}%s"
            ",p188:reply{id=R2,authorId=A1,created=2024-05-02T08:00:00+02:00}%s)"
        ) % (_txBody("one"), _txBody("two"))
        thread = ThreadedComment(_thread_element(children=children), prs_part_)

        replies = thread.replies

        assert all(isinstance(r, CommentReply) for r in replies)
        assert [(r.id, r.author, r.text) for r in replies] == [
            ("R1", "Bob", "one"),
            ("R2", "Alice", "two"),
        ]
        assert replies[0].created == dt.datetime(2024, 5, 1, 10, 0, tzinfo=dt.timezone.utc)
        assert replies[1].created == dt.datetime(
            2024, 5, 2, 8, 0, tzinfo=dt.timezone(dt.timedelta(hours=2))
        )

    def but_it_has_no_replies_when_there_are_none(self, prs_part_):
        assert ThreadedComment(_thread_element(), prs_part_).replies == []

    @pytest.mark.parametrize(("value", "expected_attr"), [(True, "resolved"), (False, None)])
    def it_can_be_resolved_and_reopened(self, value: bool, expected_attr, prs_part_):
        cm = _thread_element(",status=closed")
        thread = ThreadedComment(cm, prs_part_)

        thread.resolved = value

        assert thread.resolved is value
        assert cm.get("status") == expected_attr

    def it_can_add_a_reply(self, prs_part_):
        authors_part = prs_part_.authors_part
        prs_part_.get_or_add_authors_part.return_value = authors_part
        thread = ThreadedComment(_thread_element(children=_txBody("Fix")), prs_part_)

        reply = thread.reply("Done", "Bob")

        assert isinstance(reply, CommentReply)
        assert (reply.author, reply.text) == ("Bob", "Done")
        assert isinstance(reply.created, dt.datetime)
        assert [r.id for r in thread.replies] == [reply.id]
        assert len(authors_part) == 2

    def and_it_registers_a_new_author_for_a_reply(self, prs_part_):
        authors_part = prs_part_.authors_part
        prs_part_.get_or_add_authors_part.return_value = authors_part
        thread = ThreadedComment(_thread_element(), prs_part_)

        reply = thread.reply("Seen", "Carol", "CC")

        assert (reply.author, reply.author_initials) == ("Carol", "CC")
        assert [a.name for a in authors_part] == ["Alice", "Bob", "Carol"]


class DescribeThreadedComments:
    """Unit-test suite for `pptx.slide.ThreadedComments`.

    Reading threads from a PowerPoint-shaped file is covered in `tests/test_roundtrip.py`.
    """

    def it_is_empty_when_the_slide_has_no_modern_comments(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        assert len(slide.threaded_comments) == 0
        assert list(slide.threaded_comments) == []

    def it_raises_on_an_index_out_of_range(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        with pytest.raises(IndexError):
            slide.threaded_comments[0]
