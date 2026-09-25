"""Unit-test suite for `pptx.presentation` module."""

from __future__ import annotations

import pytest

from pptx.custom_show import CustomShows
from pptx.fonts import EmbeddedFonts
from pptx.dml.color import RGBColor
from pptx.enum.pres import PP_SLIDE_SHOW_TYPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.parts.coreprops import CorePropertiesPart
from pptx.parts.presentation import PresentationPart
from pptx.parts.presprops import PresPropsPart
from pptx.parts.slide import NotesMasterPart
from pptx.parts.tablestyles import TableStylesPart
from pptx.presentation import Presentation, Section, Sections, SlideShowSettings, TableStyles
from pptx.slide import SlideLayouts, SlideMaster, SlideMasters, Slides
from pptx.text.styles import TextListStyle
from pptx.util import Pt

from .unitutil.cxml import element, xml
from .unitutil.mock import class_mock, instance_mock, property_mock


class DescribePresentation:
    def it_provides_access_to_its_default_text_style(self):
        prs = Presentation(
            element("p:presentation/(p:sldSz,p:defaultTextStyle/(a:defPPr,a:lvl1pPr{marL=5}))"),
            None,
        )

        default_text_style = prs.default_text_style

        assert isinstance(default_text_style, TextListStyle)
        assert default_text_style[0].margin_left == 5
        assert default_text_style[1].margin_left is None

    def it_adds_a_defaultTextStyle_in_schema_order_when_written(self):
        presentation = element("p:presentation/(p:sldSz,p:notesSz,p:extLst)")
        prs = Presentation(presentation, None)

        assert prs.default_text_style[0].margin_left is None
        assert presentation.defaultTextStyle is None

        prs.default_text_style[0].font.size = Pt(18)

        assert [child.tag.split("}")[1] for child in presentation] == [
            "sldSz",
            "notesSz",
            "defaultTextStyle",
            "extLst",
        ]
        assert presentation.xpath("p:defaultTextStyle/a:lvl1pPr/a:defRPr/@sz") == ["1800"]

    def it_knows_the_height_of_its_slides(self, sld_height_get_fixture):
        prs, expected_value = sld_height_get_fixture
        assert prs.slide_height == expected_value

    def it_can_change_the_height_of_its_slides(self, sld_height_set_fixture):
        prs, slide_height, expected_xml = sld_height_set_fixture
        prs.slide_height = slide_height
        assert prs._element.xml == expected_xml

    def it_knows_the_width_of_its_slides(self, sld_width_get_fixture):
        prs, expected_value = sld_width_get_fixture
        assert prs.slide_width == expected_value

    def it_can_change_the_width_of_its_slides(self, sld_width_set_fixture):
        prs, slide_width, expected_xml = sld_width_set_fixture
        prs.slide_width = slide_width
        assert prs._element.xml == expected_xml

    def it_knows_the_height_of_its_notes_pages(self, notes_height_get_fixture):
        prs, expected_value = notes_height_get_fixture
        assert prs.notes_height == expected_value

    def it_can_change_the_height_of_its_notes_pages(self, notes_height_set_fixture):
        prs, notes_height, expected_xml = notes_height_set_fixture
        prs.notes_height = notes_height
        assert prs._element.xml == expected_xml

    def it_knows_the_width_of_its_notes_pages(self, notes_width_get_fixture):
        prs, expected_value = notes_width_get_fixture
        assert prs.notes_width == expected_value

    def it_can_change_the_width_of_its_notes_pages(self, notes_width_set_fixture):
        prs, notes_width, expected_xml = notes_width_set_fixture
        prs.notes_width = notes_width
        assert prs._element.xml == expected_xml

    def it_provides_access_to_its_custom_shows(self, custom_shows_fixture):
        prs = custom_shows_fixture
        custom_shows = prs.custom_shows
        assert isinstance(custom_shows, CustomShows)

    def it_knows_its_part(self, part_fixture):
        prs, prs_part_ = part_fixture
        assert prs.part is prs_part_

    def it_provides_access_to_its_embedded_fonts(self, prs_part_):
        prs = Presentation(
            element(
                "p:presentation/p:embeddedFontLst/p:embeddedFont/p:font{typeface=Calibri}"
            ),
            prs_part_,
        )

        embedded_fonts = prs.embedded_fonts

        assert isinstance(embedded_fonts, EmbeddedFonts)
        assert len(embedded_fonts) == 1
        assert embedded_fonts[0].typeface == "Calibri"

    def it_provides_access_to_its_core_properties(self, core_props_fixture):
        prs, core_properties_ = core_props_fixture
        assert prs.core_properties is core_properties_

    def it_provides_access_to_its_slide_show_settings(self, prs_part_):
        prs = Presentation(None, prs_part_)

        slide_show_settings = prs.slide_show_settings

        assert isinstance(slide_show_settings, SlideShowSettings)
        assert slide_show_settings._presentation_part is prs_part_

    def it_provides_access_to_its_table_styles(self, prs_part_, table_styles_part_):
        prs_part_.table_styles_part = table_styles_part_
        prs = Presentation(None, prs_part_)

        table_styles = prs.table_styles

        assert isinstance(table_styles, TableStyles)
        assert table_styles._table_styles_part is table_styles_part_

    def it_provides_access_to_its_notes_master(self, notes_master_fixture):
        prs, notes_master_ = notes_master_fixture
        assert prs.notes_master is notes_master_

    def it_provides_access_to_its_slides(self, slides_fixture):
        prs, rename_slide_parts_, rIds = slides_fixture[:3]
        Slides_, slides_, expected_xml = slides_fixture[3:]
        slides = prs.slides
        rename_slide_parts_.assert_called_once_with(rIds)
        Slides_.assert_called_once_with(prs._element.xpath("p:sldIdLst")[0], prs)
        assert prs._element.xml == expected_xml
        assert slides is slides_

    def it_provides_access_to_its_slide_layouts(self, layouts_fixture):
        prs, slide_layouts_ = layouts_fixture
        assert prs.slide_layouts is slide_layouts_

    def it_provides_access_to_its_slide_master(self, master_fixture):
        prs, getitem_, slide_master_ = master_fixture
        slide_master = prs.slide_master
        getitem_.assert_called_once_with(0)
        assert slide_master is slide_master_

    def it_provides_access_to_its_slide_masters(self, masters_fixture):
        prs, SlideMasters_, slide_masters_, expected_xml = masters_fixture
        slide_masters = prs.slide_masters
        SlideMasters_.assert_called_once_with(prs._element.xpath("p:sldMasterIdLst")[0], prs)
        assert slide_masters is slide_masters_
        assert prs._element.xml == expected_xml

    def it_can_save_the_presentation_to_a_file(self, save_fixture):
        prs, file_, prs_part_ = save_fixture
        prs.save(file_)
        prs_part_.save.assert_called_once_with(file_)

    @pytest.mark.parametrize(
        ("prs_cxml", "expected_value"),
        [
            ("p:presentation", 1),
            ("p:presentation{firstSlideNum=0}", 0),
            ("p:presentation{firstSlideNum=5}", 5),
        ],
    )
    def it_knows_its_first_slide_number(self, prs_cxml: str, expected_value: int):
        prs = Presentation(element(prs_cxml), None)
        assert prs.first_slide_number == expected_value

    @pytest.mark.parametrize(
        ("prs_cxml", "new_value", "expected_cxml"),
        [
            ("p:presentation", 0, "p:presentation{firstSlideNum=0}"),
            ("p:presentation{firstSlideNum=5}", 1, "p:presentation{firstSlideNum=1}"),
        ],
    )
    def it_can_change_its_first_slide_number(
        self, prs_cxml: str, new_value: int, expected_cxml: str
    ):
        prs = Presentation(element(prs_cxml), None)
        prs.first_slide_number = new_value
        assert prs._element.xml == xml(expected_cxml)

    # fixtures -------------------------------------------------------

    @pytest.fixture
    def custom_shows_fixture(self):
        return Presentation(element("p:presentation"), None)

    @pytest.fixture(params=[("p:presentation", None), ("p:presentation/p:notesSz{cy=42}", 42)])
    def notes_height_get_fixture(self, request):
        prs_cxml, expected_value = request.param
        prs = Presentation(element(prs_cxml), None)
        return prs, expected_value

    @pytest.fixture(
        params=[
            ("p:presentation", "p:presentation/p:notesSz{cy=914400}"),
            ("p:presentation/p:notesSz{cy=424242}", "p:presentation/p:notesSz{cy=914400}"),
        ]
    )
    def notes_height_set_fixture(self, request):
        prs_cxml, expected_cxml = request.param
        prs = Presentation(element(prs_cxml), None)
        expected_xml = xml(expected_cxml)
        return prs, 914400, expected_xml

    @pytest.fixture(params=[("p:presentation", None), ("p:presentation/p:notesSz{cx=42}", 42)])
    def notes_width_get_fixture(self, request):
        prs_cxml, expected_value = request.param
        prs = Presentation(element(prs_cxml), None)
        return prs, expected_value

    @pytest.fixture(
        params=[
            ("p:presentation", "p:presentation/p:notesSz{cx=914400}"),
            ("p:presentation/p:notesSz{cx=424242}", "p:presentation/p:notesSz{cx=914400}"),
        ]
    )
    def notes_width_set_fixture(self, request):
        prs_cxml, expected_cxml = request.param
        prs = Presentation(element(prs_cxml), None)
        expected_xml = xml(expected_cxml)
        return prs, 914400, expected_xml

    @pytest.fixture
    def core_props_fixture(self, prs_part_, core_properties_):
        prs = Presentation(None, prs_part_)
        prs_part_.core_properties = core_properties_
        return prs, core_properties_

    @pytest.fixture
    def layouts_fixture(self, masters_prop_, slide_layouts_):
        prs = Presentation(None, None)
        masters_prop_.return_value.__getitem__.return_value.slide_layouts = slide_layouts_
        return prs, slide_layouts_

    @pytest.fixture
    def master_fixture(self, masters_prop_, slide_master_):
        prs = Presentation(None, None)
        getitem_ = masters_prop_.return_value.__getitem__
        getitem_.return_value = slide_master_
        return prs, getitem_, slide_master_

    @pytest.fixture(
        params=[
            ("p:presentation", "p:presentation/p:sldMasterIdLst"),
            ("p:presentation/p:sldMasterIdLst", "p:presentation/p:sldMasterIdLst"),
        ]
    )
    def masters_fixture(self, request, SlideMasters_, slide_masters_):
        prs_cxml, expected_cxml = request.param
        prs = Presentation(element(prs_cxml), None)
        expected_xml = xml(expected_cxml)
        return prs, SlideMasters_, slide_masters_, expected_xml

    @pytest.fixture
    def notes_master_fixture(self, prs_part_, notes_master_):
        prs = Presentation(None, prs_part_)
        prs_part_.notes_master = notes_master_
        return prs, notes_master_

    @pytest.fixture
    def part_fixture(self, prs_part_):
        prs = Presentation(None, prs_part_)
        return prs, prs_part_

    @pytest.fixture
    def save_fixture(self, prs_part_):
        prs = Presentation(None, prs_part_)
        file_ = "foobar.docx"
        return prs, file_, prs_part_

    @pytest.fixture(params=[("p:presentation", None), ("p:presentation/p:sldSz{cy=42}", 42)])
    def sld_height_get_fixture(self, request):
        prs_cxml, expected_value = request.param
        prs = Presentation(element(prs_cxml), None)
        return prs, expected_value

    @pytest.fixture(
        params=[
            ("p:presentation", "p:presentation/p:sldSz{cy=914400}"),
            ("p:presentation/p:sldSz{cy=424242}", "p:presentation/p:sldSz{cy=914400}"),
        ]
    )
    def sld_height_set_fixture(self, request):
        prs_cxml, expected_cxml = request.param
        prs = Presentation(element(prs_cxml), None)
        expected_xml = xml(expected_cxml)
        return prs, 914400, expected_xml

    @pytest.fixture(params=[("p:presentation", None), ("p:presentation/p:sldSz{cx=42}", 42)])
    def sld_width_get_fixture(self, request):
        prs_cxml, expected_value = request.param
        prs = Presentation(element(prs_cxml), None)
        return prs, expected_value

    @pytest.fixture(
        params=[
            ("p:presentation", "p:presentation/p:sldSz{cx=914400}"),
            ("p:presentation/p:sldSz{cx=424242}", "p:presentation/p:sldSz{cx=914400}"),
        ]
    )
    def sld_width_set_fixture(self, request):
        prs_cxml, expected_cxml = request.param
        prs = Presentation(element(prs_cxml), None)
        expected_xml = xml(expected_cxml)
        return prs, 914400, expected_xml

    @pytest.fixture(
        params=[
            ("p:presentation", [], "p:presentation/p:sldIdLst"),
            (
                "p:presentation/p:sldIdLst/p:sldId{r:id=a}",
                ["a"],
                "p:presentation/p:sldIdLst/p:sldId{r:id=a}",
            ),
            (
                "p:presentation/p:sldIdLst/(p:sldId{r:id=a},p:sldId{r:id=b})",
                ["a", "b"],
                "p:presentation/p:sldIdLst/(p:sldId{r:id=a},p:sldId{r:id=b})",
            ),
        ]
    )
    def slides_fixture(self, request, part_prop_, Slides_, slides_):
        prs_cxml, rIds, expected_cxml = request.param
        prs = Presentation(element(prs_cxml), None)
        rename_slide_parts_ = part_prop_.return_value.rename_slide_parts
        expected_xml = xml(expected_cxml)
        return prs, rename_slide_parts_, rIds, Slides_, slides_, expected_xml

    # fixture components ---------------------------------------------

    @pytest.fixture
    def core_properties_(self, request):
        return instance_mock(request, CorePropertiesPart)

    @pytest.fixture
    def masters_prop_(self, request):
        return property_mock(request, Presentation, "slide_masters")

    @pytest.fixture
    def notes_master_(self, request):
        return instance_mock(request, NotesMasterPart)

    @pytest.fixture
    def part_prop_(self, request):
        return property_mock(request, Presentation, "part")

    @pytest.fixture
    def table_styles_part_(self, request):
        return instance_mock(request, TableStylesPart)

    @pytest.fixture
    def prs_part_(self, request):
        return instance_mock(request, PresentationPart)

    @pytest.fixture
    def slide_layouts_(self, request):
        return instance_mock(request, SlideLayouts)

    @pytest.fixture
    def SlideMasters_(self, request, slide_masters_):
        return class_mock(request, "pptx.presentation.SlideMasters", return_value=slide_masters_)

    @pytest.fixture
    def slide_master_(self, request):
        return instance_mock(request, SlideMaster)

    @pytest.fixture
    def slide_masters_(self, request):
        return instance_mock(request, SlideMasters)

    @pytest.fixture
    def Slides_(self, request, slides_):
        return class_mock(request, "pptx.presentation.Slides", return_value=slides_)

    @pytest.fixture
    def slides_(self, request):
        return instance_mock(request, Slides)


class DescribeSections:
    """Unit-test suite for `pptx.presentation.Sections` objects."""

    def _prs_elm_with_sections(self):
        xml_str = (
            '<p:presentation %s %s>'
            '  <p:extLst>'
            '    <p:ext uri="{521415D9-36F7-43E2-AB2F-B2CE04A55DE4}">'
            '      <p14:sectionLst>'
            '        <p14:section name="Intro">'
            '          <p14:sldId id="256"/>'
            '        </p14:section>'
            '        <p14:section name="Body"/>'
            '      </p14:sectionLst>'
            '    </p:ext>'
            '  </p:extLst>'
            '</p:presentation>'
        ) % (nsdecls("p"), nsdecls("p14"))
        return parse_xml(xml_str)

    def it_supports_len(self):
        prs_elm = self._prs_elm_with_sections()
        sections = Sections(prs_elm)
        assert len(sections) == 2

    def it_supports_len_with_no_sections(self):
        prs_elm = element("p:presentation")
        sections = Sections(prs_elm)
        assert len(sections) == 0

    def it_supports_indexed_access(self):
        prs_elm = self._prs_elm_with_sections()
        sections = Sections(prs_elm)
        assert sections[0].name == "Intro"
        assert sections[1].name == "Body"

    def it_supports_iteration(self):
        prs_elm = self._prs_elm_with_sections()
        sections = Sections(prs_elm)
        names = [s.name for s in sections]
        assert names == ["Intro", "Body"]

    def it_can_add_a_section(self):
        prs_elm = element("p:presentation")
        sections = Sections(prs_elm)
        section = sections.add("New Section")
        assert len(sections) == 1
        assert section.name == "New Section"

    def it_raises_on_index_out_of_range(self):
        prs_elm = element("p:presentation")
        sections = Sections(prs_elm)
        with pytest.raises(IndexError):
            sections[0]


class DescribeSection:
    """Unit-test suite for `pptx.presentation.Section` objects."""

    def _section_elm(self, name="Test", slide_ids=()):
        xml_parts = [
            '<p14:sectionLst %s>' % nsdecls("p14"),
            '  <p14:section name="%s">' % name,
        ]
        for sid in slide_ids:
            xml_parts.append('    <p14:sldId id="%d"/>' % sid)
        xml_parts.append('  </p14:section>')
        xml_parts.append('</p14:sectionLst>')
        sectionLst = parse_xml("".join(xml_parts))
        return sectionLst.section_lst[0]

    def it_knows_its_name(self):
        section = Section(self._section_elm("Intro"))
        assert section.name == "Intro"

    def it_can_change_its_name(self):
        section = Section(self._section_elm("Old"))
        section.name = "New"
        assert section.name == "New"

    def it_knows_its_slide_ids(self):
        section = Section(self._section_elm("Test", (256, 257)))
        assert section.slide_ids == (256, 257)

    def it_can_add_a_slide_id(self):
        section = Section(self._section_elm("Test"))
        section.add_slide_id(256)
        assert section.slide_ids == (256,)

    def it_can_remove_itself(self):
        xml_str = (
            '<p14:sectionLst %s>'
            '  <p14:section name="A"/>'
            '  <p14:section name="B"/>'
            '</p14:sectionLst>'
        ) % nsdecls("p14")
        sectionLst = parse_xml(xml_str)
        section = Section(sectionLst.section_lst[0])
        section.remove()
        assert len(sectionLst.section_lst) == 1
        assert sectionLst.section_lst[0].name == "B"


class _FakePresentationPart:
    """Minimal stand-in for `PresentationPart` exposing the presProps.xml hooks.

    Avoids mocking `PresentationPart` (an `XmlPart`), letting `SlideShowSettings`
    exercise the real `PresPropsPart` / oxml round-trip.
    """

    def __init__(self, pres_props_part: PresPropsPart | None = None):
        self._pres_props_part = pres_props_part

    @property
    def pres_props_part(self) -> PresPropsPart | None:
        return self._pres_props_part

    def get_or_add_pres_props_part(self) -> PresPropsPart:
        if self._pres_props_part is None:
            self._pres_props_part = PresPropsPart.new(None)
        return self._pres_props_part


class DescribeSlideShowSettings:
    """Unit-test suite for `pptx.presentation.SlideShowSettings` objects."""

    def it_reads_defaults_when_presProps_is_absent(self):
        settings = SlideShowSettings(_FakePresentationPart())

        assert settings.loop is False
        assert settings.show_narration is False
        assert settings.show_animation is True
        assert settings.use_timings is True
        assert settings.show_type == PP_SLIDE_SHOW_TYPE.SPEAKER
        assert settings.pen_color is None
        assert settings.slide_range is None
        assert settings.custom_show_id is None

    def it_does_not_create_presProps_on_a_read(self):
        prs_part = _FakePresentationPart()
        settings = SlideShowSettings(prs_part)

        settings.loop
        settings.show_type
        settings.pen_color
        settings.slide_range

        assert prs_part.pres_props_part is None

    def it_creates_presProps_on_first_write(self):
        prs_part = _FakePresentationPart()
        settings = SlideShowSettings(prs_part)

        settings.loop = True

        assert prs_part.pres_props_part is not None
        assert prs_part.pres_props_part._element.showPr.loop is True

    def it_round_trips_loop_and_kiosk_settings(self):
        settings = SlideShowSettings(_FakePresentationPart())

        settings.loop = True
        settings.show_type = PP_SLIDE_SHOW_TYPE.KIOSK

        assert settings.loop is True
        assert settings.show_type == PP_SLIDE_SHOW_TYPE.KIOSK

    @pytest.mark.parametrize(
        "show_type",
        [PP_SLIDE_SHOW_TYPE.SPEAKER, PP_SLIDE_SHOW_TYPE.BROWSE, PP_SLIDE_SHOW_TYPE.KIOSK],
    )
    def it_round_trips_each_show_type(self, show_type: PP_SLIDE_SHOW_TYPE):
        settings = SlideShowSettings(_FakePresentationPart())
        settings.show_type = show_type
        assert settings.show_type == show_type

    def it_round_trips_show_narration_animation_and_use_timings(self):
        settings = SlideShowSettings(_FakePresentationPart())

        settings.show_narration = True
        settings.show_animation = False
        settings.use_timings = False

        assert settings.show_narration is True
        assert settings.show_animation is False
        assert settings.use_timings is False

    def it_round_trips_pen_color(self):
        settings = SlideShowSettings(_FakePresentationPart())

        settings.pen_color = RGBColor(0xFF, 0x00, 0x00)

        assert settings.pen_color == RGBColor(0xFF, 0x00, 0x00)

        settings.pen_color = None
        assert settings.pen_color is None

    def it_round_trips_a_slide_range(self):
        settings = SlideShowSettings(_FakePresentationPart())

        settings.slide_range = (2, 5)

        assert settings.slide_range == (2, 5)
        assert settings.custom_show_id is None

    def it_round_trips_a_custom_show_id(self):
        settings = SlideShowSettings(_FakePresentationPart())

        settings.custom_show_id = 3

        assert settings.custom_show_id == 3
        assert settings.slide_range is None

    def it_clears_the_slide_range_choice_when_set_back_to_None(self):
        settings = SlideShowSettings(_FakePresentationPart())
        settings.slide_range = (2, 5)

        settings.slide_range = None

        assert settings.slide_range is None
        assert settings.custom_show_id is None

    def it_switches_from_slide_range_to_custom_show(self):
        settings = SlideShowSettings(_FakePresentationPart())
        settings.slide_range = (2, 5)

        settings.custom_show_id = 7

        assert settings.custom_show_id == 7
        assert settings.slide_range is None


class DescribeTableStyles:
    """Unit-test suite for `pptx.presentation.TableStyles` objects."""

    def _table_styles_part(self, styles: tuple[tuple[str, str], ...] = (), def_: str | None = None):
        from pptx.oxml import parse_xml

        parts = ['<a:tblStyleLst %s' % nsdecls("a")]
        if def_ is not None:
            parts.append(' def="%s"' % def_)
        parts.append(">")
        for styleId, styleName in styles:
            parts.append('<a:tblStyle styleId="%s" styleName="%s"/>' % (styleId, styleName))
        parts.append("</a:tblStyleLst>")
        element = parse_xml("".join(parts))
        return TableStylesPart(
            "/ppt/tableStyles.xml", "application/test", None, element=element
        )

    def it_is_empty_when_the_part_is_absent(self):
        table_styles = TableStyles(None)

        assert len(table_styles) == 0
        assert dict(table_styles) == {}
        assert table_styles.default_id is None
        assert "any-id" not in table_styles

    def it_lists_the_defined_styles(self):
        part = self._table_styles_part(
            styles=(("guid-1", "Light Style 1"), ("guid-2", "Medium Style 2")),
            def_="guid-1",
        )
        table_styles = TableStyles(part)

        assert dict(table_styles) == {"guid-1": "Light Style 1", "guid-2": "Medium Style 2"}
        assert table_styles.default_id == "guid-1"
        assert len(table_styles) == 2
        assert table_styles["guid-2"] == "Medium Style 2"

    def it_raises_KeyError_for_an_unknown_style_id(self):
        table_styles = TableStyles(self._table_styles_part())
        with pytest.raises(KeyError):
            table_styles["not-there"]
