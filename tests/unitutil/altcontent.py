"""Builds a `.pptx` whose slide holds shapes wrapped in `mc:AlternateContent`, as PowerPoint saves
them.

The slide has, in order: a text box; a 3D model (`am3d` choice, picture fallback); a slide zoom
(`pslz` choice, picture fallback); an equation (`a14` choice holding Office Math, picture-filled
shape fallback); and a shape using a 2010 drawing extension with no fallback. PowerPoint's
selection pane lists five shapes for it.
"""

from __future__ import annotations

import zipfile
from io import BytesIO

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
AM3D = "http://schemas.microsoft.com/office/drawing/2017/model3d"
PSLZ = "http://schemas.microsoft.com/office/powerpoint/2016/slidezoom"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _nv_pic(id_: int, name: str) -> str:
    return (
        '<p:nvPicPr><p:cNvPr id="%d" name="%s"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>' % (id_, name)
    )


def _pic(id_: int, name: str, x: int, y: int, cx: int, cy: int) -> str:
    return (
        "<p:pic>%s<p:blipFill><a:blip/><a:stretch><a:fillRect/></a:stretch></p:blipFill>"
        '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
        % (_nv_pic(id_, name), x, y, cx, cy)
    )


def _frame(id_: int, name: str, uri: str, payload: str, x: int, y: int, cx: int, cy: int) -> str:
    return (
        '<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="%d" name="%s"/><p:cNvGraphicFramePr/>'
        '<p:nvPr/></p:nvGraphicFramePr><p:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/>'
        '</p:xfrm><a:graphic><a:graphicData uri="%s">%s</a:graphicData></a:graphic>'
        "</p:graphicFrame>" % (id_, name, x, y, cx, cy, uri, payload)
    )


TEXTBOX_XML = (
    '<p:sp><p:nvSpPr><p:cNvPr id="10" name="TextBox 9"/><p:cNvSpPr txBox="1"/><p:nvPr/>'
    '</p:nvSpPr><p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1000" cy="1000"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr/>'
    "<a:lstStyle/><a:p><a:r><a:t>Hello</a:t></a:r></a:p></p:txBody></p:sp>"
)

MODEL3D_XML = (
    '<mc:AlternateContent xmlns:mc="%s"><mc:Choice xmlns:am3d="%s" Requires="am3d">' % (MC, AM3D)
    + _frame(
        11,
        "3D Model 10",
        AM3D,
        '<am3d:model3d r:embed="rId999"><am3d:spPr/></am3d:model3d>',
        100,
        200,
        3000,
        2000,
    )
    + "</mc:Choice><mc:Fallback>"
    + _pic(11, "3D Model 10", 110, 210, 2990, 1990)
    + "</mc:Fallback></mc:AlternateContent>"
)

ZOOM_XML = (
    '<mc:AlternateContent xmlns:mc="%s"><mc:Choice xmlns:pslz="%s" Requires="pslz">' % (MC, PSLZ)
    + _frame(
        12,
        "Slide Zoom 11",
        PSLZ,
        '<pslz:sldZm><pslz:sldZmObj sldId="257" cId="0"/></pslz:sldZm>',
        4000,
        200,
        2000,
        1125,
    )
    + "</mc:Choice><mc:Fallback>"
    + _pic(12, "Slide Zoom 11", 4000, 200, 2000, 1125)
    + "</mc:Fallback></mc:AlternateContent>"
)

EQUATION_XML = (
    '<mc:AlternateContent xmlns:mc="%s"><mc:Choice xmlns:a14="%s" Requires="a14">' % (MC, A14)
    + '<p:sp><p:nvSpPr><p:cNvPr id="13" name="TextBox 12"/><p:cNvSpPr txBox="1"/><p:nvPr/>'
    '</p:nvSpPr><p:spPr><a:xfrm><a:off x="500" y="3000"/><a:ext cx="2500" cy="400"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/>'
    '<a:p><a14:m><m:oMathPara xmlns:m="%s"><m:oMath><m:r><m:t>x=1</m:t></m:r></m:oMath>'
    "</m:oMathPara></a14:m></a:p></p:txBody></p:sp>" % M
    + "</mc:Choice><mc:Fallback>"
    '<p:sp><p:nvSpPr><p:cNvPr id="13" name="TextBox 12"/><p:cNvSpPr txBox="1"/><p:nvPr/>'
    '</p:nvSpPr><p:spPr><a:xfrm><a:off x="500" y="3000"/><a:ext cx="2500" cy="400"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:blipFill><a:blip/><a:stretch>'
    "<a:fillRect/></a:stretch></a:blipFill></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/>"
    '<a:p><a:r><a:rPr lang="en-US"><a:noFill/></a:rPr><a:t> </a:t></a:r></a:p></p:txBody></p:sp>'
    "</mc:Fallback></mc:AlternateContent>"
)

UNKNOWN_XML = (
    '<mc:AlternateContent xmlns:mc="%s"><mc:Choice xmlns:a14="%s" Requires="a14">' % (MC, A14)
    + '<p:sp><p:nvSpPr><p:cNvPr id="14" name="Rectangle 13"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="7000" y="7000"/><a:ext cx="500" cy="600"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:sp>'
    "</mc:Choice></mc:AlternateContent>"
)

SHAPES_XML = TEXTBOX_XML + MODEL3D_XML + ZOOM_XML + EQUATION_XML + UNKNOWN_XML


def altcontent_pptx() -> BytesIO:
    """Return a stream holding a one-slide `.pptx` whose slide has the five shapes above."""
    from pptx import Presentation

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    base = BytesIO()
    prs.save(base)

    out = BytesIO()
    with zipfile.ZipFile(base) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "ppt/slides/slide1.xml":
                data = (
                    data.decode("utf-8")
                    .replace("</p:spTree>", SHAPES_XML + "</p:spTree>")
                    .encode("utf-8")
                )
            dst.writestr(item, data)
    out.seek(0)
    return out
