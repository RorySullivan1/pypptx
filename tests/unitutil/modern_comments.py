"""Builds a `.pptx` package holding modern (threaded) comments, the way PowerPoint 365 saves them.

The package starts as a one-slide pypptx presentation; the authors part, the slide's modern
comments part, their relationships, content-type overrides and the slide's `p188:commentRel`
extension are then written in at the ZIP level, so loading it exercises the same path a file
from PowerPoint does.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO

P188 = "http://schemas.microsoft.com/office/powerpoint/2018/8/main"
PC = "http://schemas.microsoft.com/office/powerpoint/2013/main/command"
AC = "http://schemas.microsoft.com/office/drawing/2013/main/command"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

ALICE_ID = "{7E6A3F2B-8C1D-4E5F-9A0B-1C2D3E4F5A6B}"
BOB_ID = "{0F1E2D3C-4B5A-4978-8695-A4B3C2D1E0F9}"
THREAD_1_ID = "{11111111-2222-4333-8444-555555555555}"
THREAD_2_ID = "{66666666-7777-4888-8999-AAAAAAAAAAAA}"

SLIDE_ID = 256
CREATION_ID = 2533874937
COMMENTS_PARTNAME = "ppt/comments/modernComment_%X_%X.xml" % (SLIDE_ID, CREATION_ID)

AUTHORS_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<p188:authorLst xmlns:a="%s" xmlns:r="%s" xmlns:p188="%s">'
    '<p188:author id="%s" name="Alice Adams" initials="AA" userId="Alice Adams"'
    ' providerId="None"/>'
    '<p188:author id="%s" name="Bob Brown" initials="BB" userId="bob@example.com"'
    ' providerId="AD"/>'
    "</p188:authorLst>"
) % (A, R, P188, ALICE_ID, BOB_ID)


def _txBody(text: str) -> str:
    return (
        "<p188:txBody><a:bodyPr/><a:lstStyle/>"
        '<a:p><a:r><a:rPr lang="en-US"/><a:t>%s</a:t></a:r></a:p>'
        "</p188:txBody>" % text
    )


COMMENTS_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<p188:cmLst xmlns:a="%(a)s" xmlns:r="%(r)s" xmlns:p188="%(p188)s">'
    # -- thread 1: slide-anchored, active, two replies --
    '<p188:cm id="%(t1)s" authorId="%(alice)s" created="2024-05-01T09:30:00.123">'
    '<pc:sldMkLst xmlns:pc="%(pc)s"><pc:docMk/><pc:sldMk cId="%(cid)d" sldId="%(sid)d"/>'
    "</pc:sldMkLst>"
    '<p188:pos x="1270000" y="635000"/>'
    "<p188:replyLst>"
    '<p188:reply id="{AAAAAAAA-0000-4000-8000-000000000001}" authorId="%(bob)s"'
    ' created="2024-05-01T10:00:00.000">' + _txBody("First reply") + "</p188:reply>"
    '<p188:reply id="{AAAAAAAA-0000-4000-8000-000000000002}" authorId="%(alice)s"'
    ' created="2024-05-02T08:15:30.500Z">' + _txBody("Second reply") + "</p188:reply>"
    "</p188:replyLst>" + _txBody("Tighten this title") + "</p188:cm>"
    # -- thread 2: shape-anchored, resolved, no replies --
    '<p188:cm id="%(t2)s" authorId="%(bob)s" created="2024-05-03T14:00:00.000"'
    ' status="resolved">'
    '<ac:deMkLst xmlns:ac="%(ac)s" xmlns:pc="%(pc)s"><pc:docMk/>'
    '<pc:sldMk cId="%(cid)d" sldId="%(sid)d"/>'
    '<ac:spMk id="2" creationId="{B1B2B3B4-0000-4000-8000-C1C2C3C4C5C6}"/>'
    "</ac:deMkLst>" + _txBody("Logo is off-brand") + "</p188:cm>"
    "</p188:cmLst>"
) % {
    "a": A,
    "r": R,
    "p188": P188,
    "pc": PC,
    "ac": AC,
    "t1": THREAD_1_ID,
    "t2": THREAD_2_ID,
    "alice": ALICE_ID,
    "bob": BOB_ID,
    "cid": CREATION_ID,
    "sid": SLIDE_ID,
}

_OVERRIDES = (
    '<Override PartName="/ppt/authors.xml"'
    ' ContentType="application/vnd.ms-powerpoint.authors+xml"/>'
    '<Override PartName="/%s" ContentType="application/vnd.ms-powerpoint.comments+xml"/>'
    % COMMENTS_PARTNAME
)
_AUTHORS_REL = (
    '<Relationship Id="rId900" Type="http://schemas.microsoft.com/office/2018/10/relationships/'
    'authors" Target="authors.xml"/>'
)
_COMMENTS_REL = (
    '<Relationship Id="rId900" Type="http://schemas.microsoft.com/office/2018/10/relationships/'
    'comments" Target="../comments/%s"/>' % COMMENTS_PARTNAME.rsplit("/", 1)[1]
)
_CREATION_ID_EXT = (
    '<p:extLst><p:ext uri="{BB962C8B-B14F-4D97-AF65-F5344CB8AC3E}">'
    '<p14:creationId xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main"'
    ' val="%d"/></p:ext></p:extLst>' % CREATION_ID
)
_COMMENT_REL_EXT = (
    '<p:extLst><p:ext uri="{6950BFC3-D8DA-4A85-94F7-54DA5524770B}">'
    '<p188:commentRel xmlns:p188="%s" r:id="rId900"/></p:ext></p:extLst>' % P188
)


def modern_comments_pptx() -> BytesIO:
    """Return a stream holding a one-slide `.pptx` whose slide has two modern comment threads.

    Thread 1 (Alice, slide-anchored, active) has replies from Bob then Alice. Thread 2 (Bob,
    anchored to a shape) is resolved and has no replies.
    """
    from pptx import Presentation

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    base = BytesIO()
    prs.save(base)

    out = BytesIO()
    with zipfile.ZipFile(base) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename).decode("utf-8") if item.filename.endswith(
                (".xml", ".rels")
            ) else None
            if item.filename == "[Content_Types].xml":
                data = data.replace("</Types>", _OVERRIDES + "</Types>")
            elif item.filename == "ppt/_rels/presentation.xml.rels":
                data = data.replace("</Relationships>", _AUTHORS_REL + "</Relationships>")
            elif item.filename == "ppt/slides/_rels/slide1.xml.rels":
                data = data.replace("</Relationships>", _COMMENTS_REL + "</Relationships>")
            elif item.filename == "ppt/slides/slide1.xml":
                data = data.replace("</p:cSld>", _CREATION_ID_EXT + "</p:cSld>")
                data = re.sub(r"(</p:sld>)$", _COMMENT_REL_EXT + r"\1", data.rstrip())
            if data is None:
                dst.writestr(item, src.read(item.filename))
            else:
                dst.writestr(item, data.encode("utf-8"))
        dst.writestr("ppt/authors.xml", AUTHORS_XML)
        dst.writestr(COMMENTS_PARTNAME, COMMENTS_XML)
    out.seek(0)
    return out
