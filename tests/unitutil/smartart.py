"""Builds a `.pptx` package holding one SmartArt graphic, laid out the way PowerPoint saves it.

The package starts as a one-slide pypptx presentation; the SmartArt graphic frame, its five
diagram parts (data, layout, quick-style, colors and PowerPoint's cached drawing), their
relationships and content-type overrides are then written in at the ZIP level, so loading it
exercises the same path a file from PowerPoint does.

The node tree, in text-pane order::

    Plan            (level 0)
        Scope       (level 1)
        Budget      (level 1)
    Build           (level 0)
    <empty>         (level 0, placeholder)

Points and connections are deliberately written out of order so tree order must come from the
connections' `srcOrd`, not from document order. Presentation and transition points are present
too, as in a real file, and must not show up as nodes.
"""

from __future__ import annotations

import zipfile
from io import BytesIO

DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
DSP = "http://schemas.microsoft.com/office/drawing/2008/diagram"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

DOC_ID = "{D0C00000-0000-4000-8000-000000000000}"
PLAN_ID = "{A0000000-0000-4000-8000-000000000001}"
SCOPE_ID = "{A0000000-0000-4000-8000-000000000011}"
BUDGET_ID = "{A0000000-0000-4000-8000-000000000012}"
BUILD_ID = "{B0000000-0000-4000-8000-000000000002}"
EMPTY_ID = "{C0000000-0000-4000-8000-000000000003}"

DATA_PARTNAME = "ppt/diagrams/data1.xml"
DRAWING_PARTNAME = "ppt/diagrams/drawing1.xml"
LAYOUT_PARTNAME = "ppt/diagrams/layout1.xml"
STYLE_PARTNAME = "ppt/diagrams/quickStyle1.xml"
COLORS_PARTNAME = "ppt/diagrams/colors1.xml"

DRAWING_RID = "rId905"

_XML_DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'


def _t(text: str) -> str:
    return (
        '<dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>%s</a:t></a:r>'
        "</a:p></dgm:t>" % text
    )


def _node(model_id: str, text: str) -> str:
    return '<dgm:pt modelId="%s"><dgm:prSet phldrT="[Text]"/><dgm:spPr/>%s</dgm:pt>' % (
        model_id,
        _t(text),
    )


def _pres(model_id: str, assoc_id: str) -> str:
    return (
        '<dgm:pt modelId="%s" type="pres"><dgm:prSet presAssocID="%s" presName="node"'
        ' presStyleLbl="node1" presStyleIdx="0" presStyleCnt="3"/><dgm:spPr/></dgm:pt>'
        % (model_id, assoc_id)
    )


def _cxn(model_id: str, src: str, dest: str, ord_: int, type_: str | None = None) -> str:
    type_attr = ' type="%s"' % type_ if type_ else ""
    return '<dgm:cxn modelId="%s"%s srcId="%s" destId="%s" srcOrd="%d" destOrd="0"/>' % (
        model_id,
        type_attr,
        src,
        dest,
        ord_,
    )


DATA_XML = (
    _XML_DECL + '<dgm:dataModel xmlns:dgm="%s" xmlns:a="%s" xmlns:r="%s">' % (DGM, A, R)
    + "<dgm:ptLst>"
    + '<dgm:pt modelId="%s" type="doc"><dgm:prSet loTypeId="urn:microsoft.com/office/officeart/'
    '2005/8/layout/default"/><dgm:spPr/><dgm:t><a:bodyPr/><a:lstStyle/>'
    '<a:p><a:endParaRPr lang="en-US"/></a:p></dgm:t></dgm:pt>' % DOC_ID
    + '<dgm:pt modelId="%s"><dgm:prSet phldr="1"/><dgm:spPr/></dgm:pt>' % EMPTY_ID
    + _node(BUDGET_ID, "Budget")
    + _node(PLAN_ID, "Plan")
    + '<dgm:pt modelId="{E0000000-0000-4000-8000-000000000001}" type="parTrans"'
    ' cxnId="{F0000000-0000-4000-8000-000000000001}"><dgm:prSet/><dgm:spPr/>'
    "<dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr/></a:p></dgm:t></dgm:pt>"
    + '<dgm:pt modelId="{E0000000-0000-4000-8000-000000000002}" type="sibTrans"'
    ' cxnId="{F0000000-0000-4000-8000-000000000001}"><dgm:prSet/><dgm:spPr/>'
    "<dgm:t><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr/></a:p></dgm:t></dgm:pt>"
    + _node(BUILD_ID, "Build")
    + _node(SCOPE_ID, "Scope")
    + _pres("{90000000-0000-4000-8000-000000000001}", PLAN_ID)
    + _pres("{90000000-0000-4000-8000-000000000002}", BUILD_ID)
    + "</dgm:ptLst>"
    + "<dgm:cxnLst>"
    + _cxn("{F0000000-0000-4000-8000-000000000005}", PLAN_ID, BUDGET_ID, 1)
    + _cxn("{F0000000-0000-4000-8000-000000000003}", DOC_ID, EMPTY_ID, 2)
    + _cxn("{F0000000-0000-4000-8000-000000000004}", PLAN_ID, SCOPE_ID, 0, "parOf")
    + _cxn("{F0000000-0000-4000-8000-000000000002}", DOC_ID, BUILD_ID, 1)
    + _cxn("{F0000000-0000-4000-8000-000000000001}", DOC_ID, PLAN_ID, 0)
    + _cxn(
        "{F0000000-0000-4000-8000-000000000006}",
        PLAN_ID,
        "{90000000-0000-4000-8000-000000000001}",
        0,
        "presOf",
    )
    + _cxn(
        "{F0000000-0000-4000-8000-000000000007}",
        "{90000000-0000-4000-8000-000000000001}",
        "{90000000-0000-4000-8000-000000000002}",
        0,
        "presParOf",
    )
    + "</dgm:cxnLst>"
    + "<dgm:bg/><dgm:whole/>"
    + '<dgm:extLst><a:ext uri="http://schemas.microsoft.com/office/drawing/2008/diagram">'
    '<dsp:dataModelExt xmlns:dsp="%s" relId="%s" minVer="http://schemas.openxmlformats.org/'
    'drawingml/2006/diagram"/></a:ext></dgm:extLst>' % (DSP, DRAWING_RID)
    + "</dgm:dataModel>"
)


def _dsp_sp(model_id: str, text: str, y: int) -> str:
    return (
        '<dsp:sp modelId="%s"><dsp:nvSpPr><dsp:cNvPr id="0" name=""/><dsp:cNvSpPr/>'
        '</dsp:nvSpPr><dsp:spPr><a:xfrm><a:off x="0" y="%d"/><a:ext cx="3000000" cy="600000"/>'
        '</a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></dsp:spPr>'
        '<dsp:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>%s</a:t>'
        "</a:r></a:p></dsp:txBody></dsp:sp>" % (model_id, y, text)
    )


DRAWING_XML = (
    _XML_DECL
    + '<dsp:drawing xmlns:dgm="%s" xmlns:dsp="%s" xmlns:a="%s">' % (DGM, DSP, A)
    + '<dsp:spTree><dsp:nvGrpSpPr><dsp:cNvPr id="0" name=""/><dsp:cNvGrpSpPr/></dsp:nvGrpSpPr>'
    "<dsp:grpSpPr/>"
    + _dsp_sp("{90000000-0000-4000-8000-000000000001}", "Plan", 0)
    + _dsp_sp("{90000000-0000-4000-8000-000000000002}", "Build", 700000)
    + "</dsp:spTree></dsp:drawing>"
)

LAYOUT_XML = (
    _XML_DECL + '<dgm:layoutDef xmlns:dgm="%s" xmlns:a="%s"'
    ' uniqueId="urn:microsoft.com/office/officeart/2005/8/layout/default">'
    '<dgm:title val=""/><dgm:desc val=""/><dgm:layoutNode name="diagram"/></dgm:layoutDef>'
    % (DGM, A)
)
STYLE_XML = (
    _XML_DECL + '<dgm:styleDef xmlns:dgm="%s" xmlns:a="%s"'
    ' uniqueId="urn:microsoft.com/office/officeart/2005/8/quickstyle/simple1">'
    '<dgm:title val=""/><dgm:desc val=""/></dgm:styleDef>' % (DGM, A)
)
COLORS_XML = (
    _XML_DECL + '<dgm:colorsDef xmlns:dgm="%s" xmlns:a="%s"'
    ' uniqueId="urn:microsoft.com/office/officeart/2005/8/colors/accent1_2">'
    '<dgm:title val=""/><dgm:desc val=""/></dgm:colorsDef>' % (DGM, A)
)

GRAPHIC_FRAME_XML = (
    "<p:graphicFrame><p:nvGraphicFramePr>"
    '<p:cNvPr id="42" name="Diagram 41"/><p:cNvGraphicFramePr/><p:nvPr/>'
    "</p:nvGraphicFramePr>"
    '<p:xfrm><a:off x="914400" y="914400"/><a:ext cx="6096000" cy="4064000"/></p:xfrm>'
    '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">'
    '<dgm:relIds xmlns:dgm="%s" r:dm="rId901" r:lo="rId902" r:qs="rId903" r:cs="rId904"/>'
    "</a:graphicData></a:graphic></p:graphicFrame>" % DGM
)

_OVERRIDES = "".join(
    '<Override PartName="/%s" ContentType="application/%s"/>' % (partname, content_type)
    for partname, content_type in (
        (DATA_PARTNAME, "vnd.openxmlformats-officedocument.drawingml.diagramData+xml"),
        (LAYOUT_PARTNAME, "vnd.openxmlformats-officedocument.drawingml.diagramLayout+xml"),
        (STYLE_PARTNAME, "vnd.openxmlformats-officedocument.drawingml.diagramStyle+xml"),
        (COLORS_PARTNAME, "vnd.openxmlformats-officedocument.drawingml.diagramColors+xml"),
        (DRAWING_PARTNAME, "vnd.ms-office.drawingml.diagramDrawing+xml"),
    )
)
_SLIDE_RELS = "".join(
    '<Relationship Id="%s" Type="%s" Target="../diagrams/%s"/>'
    % (rId, reltype, partname.rsplit("/", 1)[1])
    for rId, reltype, partname in (
        ("rId901", R + "/diagramData", DATA_PARTNAME),
        ("rId902", R + "/diagramLayout", LAYOUT_PARTNAME),
        ("rId903", R + "/diagramQuickStyle", STYLE_PARTNAME),
        ("rId904", R + "/diagramColors", COLORS_PARTNAME),
        (
            DRAWING_RID,
            "http://schemas.microsoft.com/office/2007/relationships/diagramDrawing",
            DRAWING_PARTNAME,
        ),
    )
)

DIAGRAM_PARTS = {
    DATA_PARTNAME: DATA_XML,
    LAYOUT_PARTNAME: LAYOUT_XML,
    STYLE_PARTNAME: STYLE_XML,
    COLORS_PARTNAME: COLORS_XML,
    DRAWING_PARTNAME: DRAWING_XML,
}


def smartart_pptx() -> BytesIO:
    """Return a stream holding a one-slide `.pptx` whose slide has one SmartArt graphic frame."""
    from pptx import Presentation

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    base = BytesIO()
    prs.save(base)

    out = BytesIO()
    with zipfile.ZipFile(base) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = (
                src.read(item.filename).decode("utf-8")
                if item.filename.endswith((".xml", ".rels"))
                else None
            )
            if item.filename == "[Content_Types].xml":
                data = data.replace("</Types>", _OVERRIDES + "</Types>")
            elif item.filename == "ppt/slides/_rels/slide1.xml.rels":
                data = data.replace("</Relationships>", _SLIDE_RELS + "</Relationships>")
            elif item.filename == "ppt/slides/slide1.xml":
                data = data.replace("</p:spTree>", GRAPHIC_FRAME_XML + "</p:spTree>")
            if data is None:
                dst.writestr(item, src.read(item.filename))
            else:
                dst.writestr(item, data.encode("utf-8"))
        for partname, xml in DIAGRAM_PARTS.items():
            dst.writestr(partname, xml.encode("utf-8"))
    out.seek(0)
    return out
