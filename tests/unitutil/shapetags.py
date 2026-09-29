"""Builds a `.pptx` whose slide and shapes carry customer data tags, as PowerPoint saves them.

The slide has a tagged rectangle, an untagged rectangle and a group whose one child is tagged,
and the slide itself has tags. As PowerPoint writes it, each tagged shape's `p:nvPr` and the
slide's `p:cSld` hold a `p:custDataLst/p:tags` reference to a tags part of their own
(`/ppt/tags/tagN.xml`), related from the slide part. PowerPoint writes tag names in upper case.
"""

from __future__ import annotations

import zipfile
from io import BytesIO

TAGS_CT = "application/vnd.openxmlformats-officedocument.presentationml.tags+xml"
TAGS_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/tags"


def _sp(id_: int, name: str, nvPr: str, x: int) -> str:
    return (
        '<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s"/><p:cNvSpPr/>%s</p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="%d" y="0"/><a:ext cx="1000" cy="1000"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:sp>' % (id_, name, nvPr, x)
    )


def _tagged_nvPr(rId: str) -> str:
    return '<p:nvPr><p:custDataLst><p:tags r:id="%s"/></p:custDataLst></p:nvPr>' % rId


SHAPES_XML = (
    _sp(10, "Tagged 9", _tagged_nvPr("rId90"), 0)
    + _sp(11, "Plain 10", "<p:nvPr/>", 2000)
    + '<p:grpSp><p:nvGrpSpPr><p:cNvPr id="12" name="Group 11"/><p:cNvGrpSpPr/><p:nvPr/>'
    "</p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x=\"4000\" y=\"0\"/><a:ext cx=\"1000\" cy=\"1000\"/>"
    '<a:chOff x="4000" y="0"/><a:chExt cx="1000" cy="1000"/></a:xfrm></p:grpSpPr>'
    + _sp(13, "Child 12", _tagged_nvPr("rId91"), 4000)
    + "</p:grpSp>"
)

SLIDE_CUSTDATA_XML = '<p:custDataLst><p:tags r:id="rId92"/></p:custDataLst>'

TAG_PARTS = {
    "rId90": ("tag1.xml", {"SOURCE": "crm", "GENERATED": "1"}),
    "rId91": ("tag2.xml", {"ROLE": "child"}),
    "rId92": ("tag3.xml", {"STATUS": "draft"}),
}


def _tagLst_xml(tags: dict[str, str]) -> bytes:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<p:tagLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        + "".join('<p:tag name="%s" val="%s"/>' % item for item in tags.items())
        + "</p:tagLst>"
    ).encode("utf-8")


def shape_tags_pptx() -> BytesIO:
    """Return a stream holding a one-slide `.pptx` with the tagged slide and shapes above."""
    from pptx import Presentation

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    base = BytesIO()
    prs.save(base)

    rels_xml = "".join(
        '<Relationship Id="%s" Type="%s" Target="../tags/%s"/>' % (rId, TAGS_RT, filename)
        for rId, (filename, _) in TAG_PARTS.items()
    )
    overrides_xml = "".join(
        '<Override PartName="/ppt/tags/%s" ContentType="%s"/>' % (filename, TAGS_CT)
        for filename, _ in TAG_PARTS.values()
    )

    edits = {
        "ppt/slides/slide1.xml": ("</p:spTree>", SHAPES_XML + "</p:spTree>" + SLIDE_CUSTDATA_XML),
        "ppt/slides/_rels/slide1.xml.rels": ("</Relationships>", rels_xml + "</Relationships>"),
        "[Content_Types].xml": ("</Types>", overrides_xml + "</Types>"),
    }

    out = BytesIO()
    with zipfile.ZipFile(base) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename in edits:
                old, new = edits[item.filename]
                data = data.decode("utf-8").replace(old, new).encode("utf-8")
            dst.writestr(item, data)
        for filename, tags in TAG_PARTS.values():
            dst.writestr("ppt/tags/%s" % filename, _tagLst_xml(tags))
    out.seek(0)
    return out
