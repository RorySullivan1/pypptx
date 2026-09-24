"""Integration test suite for embedded-fonts support (issue #61).

Builds a real `.pptx` file containing an embedded font part (using the
`tests/test_files/calibriz.ttf` fixture), writes it to disk with the actual OPC
package writer, and confirms it can be listed and removed via the public API, with
the on-disk zip left free of dangling relationships / content-type overrides after
a font is removed.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

test_files_dir = Path(__file__).parent / "test_files"
calibriz_ttf_path = test_files_dir / "calibriz.ttf"


def _build_presentation_with_embedded_font(typeface: str = "Calibri"):
    """Return a `Presentation` with a `typeface` font part related and listed."""
    prs = Presentation()
    prs_part = prs.part

    font_bytes = calibriz_ttf_path.read_bytes()
    font_part = Part(
        PackURI("/ppt/fonts/font1.fntdata"), CT.X_FONTDATA, prs_part.package, font_bytes
    )
    rId = prs_part.relate_to(font_part, RT.FONT)

    embeddedFontLst = prs._element.get_or_add_embeddedFontLst()
    entry = embeddedFontLst.add_embeddedFont(typeface)
    entry.add_style("regular", rId)

    return prs


class DescribeEmbeddedFontsIntegration:
    """Integration tests exercising real save/load of an embedded-fonts pptx."""

    def it_lists_an_embedded_font_after_a_real_save_and_load_cycle(self, tmp_path):
        prs = _build_presentation_with_embedded_font("Calibri")
        pptx_path = tmp_path / "with-font.pptx"
        prs.save(str(pptx_path))

        # -- the font part and its relationship are really in the zip --
        with zipfile.ZipFile(pptx_path) as zf:
            names = zf.namelist()
        assert any(name.endswith("font1.fntdata") for name in names)

        prs2 = Presentation(str(pptx_path))
        embedded_fonts = prs2.embedded_fonts

        assert len(embedded_fonts) == 1
        font = embedded_fonts[0]
        assert font.typeface == "Calibri"
        assert font.styles == ("regular",)
        assert font.has_regular is True
        assert font.has_bold is False

    def it_writes_font_data_back_unchanged_when_untouched(self, tmp_path):
        first_path, second_path = tmp_path / "first.pptx", tmp_path / "second.pptx"
        _build_presentation_with_embedded_font().save(str(first_path))

        Presentation(str(first_path)).save(str(second_path))

        with zipfile.ZipFile(second_path) as zf:
            assert zf.read("ppt/fonts/font1.fntdata") == calibriz_ttf_path.read_bytes()

    def it_removes_the_font_part_and_relationship_on_save_after_remove(self, tmp_path):
        prs = _build_presentation_with_embedded_font("Calibri")
        pptx_path = tmp_path / "with-font.pptx"
        prs.save(str(pptx_path))

        prs2 = Presentation(str(pptx_path))
        font = prs2.embedded_fonts[0]
        font.remove()

        reopened_path = tmp_path / "font-removed.pptx"
        prs2.save(str(reopened_path))

        with zipfile.ZipFile(reopened_path) as zf:
            names = zf.namelist()
            content_types_xml = zf.read("[Content_Types].xml").decode("utf-8")
            prs_rels_xml = zf.read("ppt/_rels/presentation.xml.rels").decode("utf-8")
            prs_xml = zf.read("ppt/presentation.xml").decode("utf-8")

        # -- the font part itself is gone --
        assert not any(name.endswith(".fntdata") for name in names)
        # -- no dangling relationship to it remains --
        assert "font1.fntdata" not in prs_rels_xml
        assert RT.FONT not in prs_rels_xml
        # -- no stale content-type override for the font part remains --
        assert "fntdata" not in content_types_xml
        # -- the (now empty) embeddedFontLst was dropped and re-embedding turned off --
        assert "embeddedFontLst" not in prs_xml
        assert 'embedTrueTypeFonts="0"' in prs_xml

        # -- and the file still opens cleanly with no embedded fonts reported --
        prs3 = Presentation(str(reopened_path))
        assert len(prs3.embedded_fonts) == 0
