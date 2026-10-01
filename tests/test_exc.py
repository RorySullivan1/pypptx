"""Unit-test suite for `pptx.exc`: the exception hierarchy callers rely on."""

from __future__ import annotations

import pytest

from pptx.exc import (
    ChartError,
    InvalidPackageError,
    InvalidTypeError,
    InvalidValueError,
    PackageError,
    PyPptxError,
    ShapeError,
    SlideError,
    TableError,
)
from pptx.util import Inches


class DescribeExceptionHierarchy:
    """Each pypptx exception keeps the builtin base the API raised before it existed."""

    @pytest.mark.parametrize(
        ("exc_cls", "builtin"),
        [
            (InvalidValueError, ValueError),
            (InvalidTypeError, TypeError),
            (ShapeError, ValueError),
            (SlideError, ValueError),
            (ChartError, ValueError),
            (TableError, ValueError),
        ],
    )
    def it_is_both_a_PyPptxError_and_the_builtin_it_replaced(self, exc_cls, builtin):
        assert issubclass(exc_cls, PyPptxError)
        assert issubclass(exc_cls, builtin)

    def it_reports_an_unreadable_package_as_a_PackageError(self):
        assert issubclass(InvalidPackageError, PackageError)
        assert not issubclass(InvalidPackageError, ValueError)


class DescribeValidationErrors:
    """Property setters reject bad values with the new classes, still catchable as builtins."""

    def it_raises_InvalidValueError_for_an_out_of_range_value(self):
        from pptx import Presentation

        prs = Presentation()
        textbox = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
            0, 0, Inches(1), Inches(1)
        )
        font = textbox.text_frame.paragraphs[0].font

        with pytest.raises(InvalidValueError):
            font.size = -1
        with pytest.raises(ValueError):
            font.size = -1

    def it_raises_InvalidTypeError_for_a_value_of_the_wrong_type(self):
        from pptx import Presentation

        prs = Presentation()
        textbox = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
            0, 0, Inches(1), Inches(1)
        )

        with pytest.raises(InvalidTypeError):
            textbox.left = "one inch"
        with pytest.raises(TypeError):
            textbox.left = "one inch"
