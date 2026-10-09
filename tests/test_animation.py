"""Unit-test suite for the `pptx.animation` module."""

from __future__ import annotations

import datetime as dt

import pytest

from pptx.animation import Animation, SlideAnimations
from pptx.enum.animation import MSO_ANIMATION_EFFECT, MSO_ANIMATION_TRIGGER, PP_ANIMATION_CLASS
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.slide import Slide


def _sp(shape_id: int, name: str) -> str:
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{shape_id}" name="{name}"/><p:cNvSpPr/><p:nvPr/>'
        "</p:nvSpPr><p:spPr/></p:sp>"
    )


def _effect(
    cTn_id: int,
    node_type: str | None,
    spid: int,
    preset_class: str = "entr",
    preset_id: int = 10,
    delay: str = "0",
    behaviors: str | None = None,
    pRg: tuple[int, int] | None = None,
) -> str:
    """One effect `p:par` the way PowerPoint writes it."""
    node_type_attr = "" if node_type is None else f' nodeType="{node_type}"'
    txEl = "" if pRg is None else f'<p:txEl><p:pRg st="{pRg[0]}" end="{pRg[1]}"/></p:txEl>'
    tgtEl = f'<p:tgtEl><p:spTgt spid="{spid}">{txEl}</p:spTgt></p:tgtEl>'
    if behaviors is None:
        behaviors = (
            f'<p:set><p:cBhvr><p:cTn id="{cTn_id + 1}" dur="1" fill="hold"/>{tgtEl}</p:cBhvr>'
            "</p:set>"
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr>'
            f'<p:cTn id="{cTn_id + 2}" dur="500"/>{tgtEl}</p:cBhvr></p:animEffect>'
        )
    return (
        f'<p:par><p:cTn id="{cTn_id}" presetID="{preset_id}" presetClass="{preset_class}"'
        f' presetSubtype="0" fill="hold"{node_type_attr}>'
        f'<p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>'
        f"<p:childTnLst>{behaviors}</p:childTnLst></p:cTn></p:par>"
    )


def _click_group(cTn_id: int, *effects: str) -> str:
    return (
        f'<p:par><p:cTn id="{cTn_id}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
        f'</p:stCondLst><p:childTnLst><p:par><p:cTn id="{cTn_id + 1}" fill="hold">'
        '<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f"<p:childTnLst>{''.join(effects)}</p:childTnLst></p:cTn></p:par>"
        "</p:childTnLst></p:cTn></p:par>"
    )


def _main_seq(*groups: str) -> str:
    return (
        '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
        f"<p:childTnLst>{''.join(groups)}</p:childTnLst></p:cTn></p:seq>"
    )


def _interactive_seq(trigger_spid: int, evt: str, *effects: str) -> str:
    return (
        '<p:seq concurrent="1" nextAc="seek"><p:cTn id="90" restart="whenNotActive" fill="hold"'
        ' evtFilter="cancelBubble" nodeType="interactiveSeq"><p:stCondLst>'
        f'<p:cond evt="{evt}" delay="0"><p:tgtEl><p:spTgt spid="{trigger_spid}"/></p:tgtEl>'
        "</p:cond></p:stCondLst><p:childTnLst>"
        f'<p:par><p:cTn id="91" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f"<p:childTnLst>{''.join(effects)}</p:childTnLst></p:cTn></p:par>"
        "</p:childTnLst></p:cTn></p:seq>"
    )


def _slide(*sequences: str, shapes: str | None = None) -> Slide:
    if shapes is None:
        shapes = _sp(3, "Title") + _sp(4, "Body") + _sp(5, "Button")
    timing = (
        ""
        if not sequences
        else '<p:timing><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never"'
        f' nodeType="tmRoot"><p:childTnLst>{"".join(sequences)}</p:childTnLst></p:cTn>'
        "</p:par></p:tnLst></p:timing>"
    )
    sld = parse_xml(
        f"<p:sld {nsdecls('a', 'p')}><p:cSld><p:spTree><p:nvGrpSpPr><p:cNvPr id=\"1\""
        ' name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>'
        f"{shapes}</p:spTree></p:cSld>{timing}</p:sld>"
    )
    return Slide(sld, None)  # pyright: ignore[reportArgumentType]


class DescribeSlideAnimations:
    """Unit-test suite for `pptx.animation.SlideAnimations` objects."""

    def it_is_provided_by_a_slide(self):
        animations = _slide().animations

        assert isinstance(animations, SlideAnimations)
        assert len(animations) == 0
        assert list(animations) == []

    def it_lists_the_main_sequence_in_play_order_then_triggered_sequences(self):
        slide = _slide(
            _interactive_seq(5, "onClick", _effect(100, "clickEffect", spid=3, preset_id=2)),
            _main_seq(
                _click_group(10, _effect(12, "clickEffect", spid=3)),
                _click_group(20, _effect(22, "clickEffect", spid=4), _effect(30, "withEffect", 3)),
            ),
        )

        animations = slide.animations

        assert len(animations) == 4
        assert all(isinstance(a, Animation) for a in animations)
        assert [(a.shape_id, a.trigger) for a in animations] == [
            (3, MSO_ANIMATION_TRIGGER.ON_PAGE_CLICK),
            (4, MSO_ANIMATION_TRIGGER.ON_PAGE_CLICK),
            (3, MSO_ANIMATION_TRIGGER.WITH_PREVIOUS),
            (3, MSO_ANIMATION_TRIGGER.ON_SHAPE_CLICK),
        ]
        assert animations[1].shape_id == 4
        assert animations[-1].preset_id == 2

    def it_does_not_list_media_nodes(self):
        video = (
            '<p:video><p:cMediaNode vol="80000"><p:cTn id="80" fill="hold" display="0">'
            '<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn>'
            '<p:tgtEl><p:spTgt spid="5"/></p:tgtEl></p:cMediaNode></p:video>'
        )
        slide = _slide(_main_seq(_click_group(10, _effect(12, "clickEffect", 3))), video)

        assert [a.shape_id for a in slide.animations] == [3]


class DescribeAnimation:
    """Unit-test suite for `pptx.animation.Animation` objects."""

    def it_knows_its_preset(self):
        slide = _slide(_main_seq(_click_group(10, _effect(12, "clickEffect", 3, "exit", 10))))

        (animation,) = slide.animations

        assert animation.preset_class == PP_ANIMATION_CLASS.EXIT
        assert animation.preset_id == 10
        assert animation.preset_subtype == 0
        assert animation.effect_type == MSO_ANIMATION_EFFECT.FADE

    @pytest.mark.parametrize(
        ("preset_class", "preset_id"),
        [("entr", 37), ("emph", 8), ("path", 1), ("entr", 0)],
    )
    def but_it_names_only_entrance_and_exit_effects_1_to_31(self, preset_class, preset_id):
        effect = _effect(12, "clickEffect", 3, preset_class, preset_id)
        slide = _slide(_main_seq(_click_group(10, effect)))

        (animation,) = slide.animations

        assert animation.preset_id == preset_id
        assert animation.effect_type is None

    @pytest.mark.parametrize(
        ("node_type", "expected_trigger"),
        [
            ("clickEffect", MSO_ANIMATION_TRIGGER.ON_PAGE_CLICK),
            ("withEffect", MSO_ANIMATION_TRIGGER.WITH_PREVIOUS),
            ("afterEffect", MSO_ANIMATION_TRIGGER.AFTER_PREVIOUS),
            (None, MSO_ANIMATION_TRIGGER.ON_PAGE_CLICK),
        ],
    )
    def it_knows_what_triggers_it(self, node_type, expected_trigger):
        slide = _slide(_main_seq(_click_group(10, _effect(12, node_type, 3))))

        (animation,) = slide.animations

        assert animation.trigger == expected_trigger
        assert animation.trigger_shape_id is None
        assert animation.trigger_shape is None

    def it_knows_the_shape_that_triggers_a_triggered_sequence(self):
        slide = _slide(_interactive_seq(5, "onClick", _effect(100, "clickEffect", spid=3)))

        (animation,) = slide.animations

        assert animation.trigger == MSO_ANIMATION_TRIGGER.ON_SHAPE_CLICK
        assert animation.trigger_shape_id == 5
        assert animation.trigger_shape is not None
        assert animation.trigger_shape.name == "Button"

    def and_it_recognizes_a_media_bookmark_trigger(self):
        slide = _slide(_interactive_seq(5, "onMediaBookmark", _effect(100, "clickEffect", 3)))

        (animation,) = slide.animations

        assert animation.trigger == MSO_ANIMATION_TRIGGER.ON_MEDIA_BOOKMARK

    def it_knows_its_delay_and_duration(self):
        behaviors = (
            '<p:set><p:cBhvr><p:cTn id="13" dur="1" fill="hold">'
            '<p:stCondLst><p:cond delay="499"/></p:stCondLst></p:cTn>'
            '<p:tgtEl><p:spTgt spid="3"/></p:tgtEl></p:cBhvr></p:set>'
            '<p:anim><p:cBhvr><p:cTn id="14" dur="750"/>'
            '<p:tgtEl><p:spTgt spid="3"/></p:tgtEl></p:cBhvr></p:anim>'
        )
        effect = _effect(12, "clickEffect", 3, delay="250", behaviors=behaviors)
        slide = _slide(_main_seq(_click_group(10, effect)))

        (animation,) = slide.animations

        assert animation.delay == dt.timedelta(milliseconds=250)
        assert animation.duration == dt.timedelta(milliseconds=750)

    @pytest.mark.parametrize(
        ("delay", "behavior_dur", "expected_delay", "expected_duration"),
        [
            ("indefinite", "500", None, dt.timedelta(milliseconds=500)),
            ("0", "indefinite", dt.timedelta(0), None),
        ],
    )
    def and_it_reads_an_indefinite_time_as_None(
        self, delay, behavior_dur, expected_delay, expected_duration
    ):
        behaviors = (
            f'<p:anim><p:cBhvr><p:cTn id="14" dur="{behavior_dur}"/>'
            '<p:tgtEl><p:spTgt spid="3"/></p:tgtEl></p:cBhvr></p:anim>'
        )
        effect = _effect(12, "clickEffect", 3, delay=delay, behaviors=behaviors)
        slide = _slide(_main_seq(_click_group(10, effect)))

        (animation,) = slide.animations

        assert animation.delay == expected_delay
        assert animation.duration == expected_duration

    def it_knows_the_shape_and_paragraphs_it_animates(self):
        slide = _slide(_main_seq(_click_group(10, _effect(12, "clickEffect", 4, pRg=(1, 2)))))

        (animation,) = slide.animations

        assert animation.shape_id == 4
        assert animation.shape is not None
        assert animation.shape.name == "Body"
        assert animation.paragraphs == range(1, 3)

    def and_it_finds_a_shape_inside_a_group(self):
        group = (
            '<p:grpSp><p:nvGrpSpPr><p:cNvPr id="6" name="Group"/><p:cNvGrpSpPr/><p:nvPr/>'
            f"</p:nvGrpSpPr><p:grpSpPr/>{_sp(7, 'Inner')}</p:grpSp>"
        )
        slide = _slide(_main_seq(_click_group(10, _effect(12, "clickEffect", 7))), shapes=group)

        (animation,) = slide.animations

        assert animation.shape is not None
        assert animation.shape.name == "Inner"
        assert animation.paragraphs is None

    def but_it_has_no_shape_when_the_target_is_not_on_the_slide(self):
        slide = _slide(_main_seq(_click_group(10, _effect(12, "clickEffect", 42))))

        (animation,) = slide.animations

        assert animation.shape_id == 42
        assert animation.shape is None

    def it_reads_markup_that_breaks_the_schema_without_raising(self):
        effect = (
            '<p:par><p:cTn id="12" presetID="x" presetClass="sparkle" presetSubtype="?"'
            ' nodeType="clickEffect"><p:stCondLst><p:cond delay="soon"/></p:stCondLst>'
            '<p:childTnLst><p:set><p:cBhvr><p:cTn id="13" dur="-"/>'
            '<p:tgtEl><p:spTgt spid="three"><p:txEl><p:pRg st="a" end="b"/></p:txEl>'
            "</p:spTgt></p:tgtEl></p:cBhvr></p:set></p:childTnLst></p:cTn></p:par>"
        )
        slide = _slide(_main_seq(_click_group(10, effect)))

        (animation,) = slide.animations

        assert animation.preset_class is None
        assert animation.preset_id is None
        assert animation.preset_subtype is None
        assert animation.effect_type is None
        assert animation.delay == dt.timedelta(0)
        assert animation.duration is None
        assert animation.shape_id is None
        assert animation.shape is None
        assert animation.paragraphs is None
        assert animation.trigger == MSO_ANIMATION_TRIGGER.ON_PAGE_CLICK

    def it_does_not_change_the_xml_when_read(self):
        slide = _slide(
            _main_seq(_click_group(10, _effect(12, "clickEffect", 4, pRg=(0, 0)))),
            _interactive_seq(5, "onClick", _effect(100, "clickEffect", 3)),
        )
        before = slide._element.xml

        for a in slide.animations:
            _ = (a.preset_class, a.preset_id, a.preset_subtype, a.effect_type, a.trigger)
            _ = (a.trigger_shape, a.delay, a.duration, a.shape, a.paragraphs)

        assert slide._element.xml == before
