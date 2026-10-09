"""Behavior and round-trip tests for authoring slide layouts (v0.5.0).

These build on the bundled default template and the real-world corpus rather than mocks: what
matters is the package a layout leaves behind, and what a reloaded presentation reads from it.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.exc import ShapeError, SlideError
from pptx.shapes.placeholder import LayoutPlaceholder
from pptx.util import Inches

CORPUS_DIR = Path(__file__).parent / "test_files" / "real_world"


def _reloaded(prs):
    stream = io.BytesIO()
    prs.save(stream)
    stream.seek(0)
    return Presentation(stream)


def _master_and_layout_ids(prs) -> list[int]:
    ids = [e.id for e in prs.part._element.sldMasterIdLst.sldMasterId_lst]
    for master in prs.slide_masters:
        ids.extend(e.id for e in master._element.sldLayoutIdLst.sldLayoutId_lst)
    return ids


def _placeholder_summary(layout):
    return [(p.placeholder_format.type, p.placeholder_format.idx) for p in layout.placeholders]


class DescribeSlideLayouts_add_slide_layout:
    """`SlideLayouts.add_slide_layout()`."""

    def it_adds_a_layout_like_PowerPoints_insert_layout(self):
        prs = Presentation()
        master = prs.slide_masters[0]
        count = len(master.slide_layouts)

        layout = master.slide_layouts.add_slide_layout("Quote")

        assert len(master.slide_layouts) == count + 1
        assert master.slide_layouts[-1] == layout
        assert layout.name == "Quote"
        assert layout.slide_master == master
        assert _placeholder_summary(layout) == [
            (PP_PLACEHOLDER.TITLE, 0),
            (PP_PLACEHOLDER.DATE, 10),
            (PP_PLACEHOLDER.FOOTER, 11),
            (PP_PLACEHOLDER.SLIDE_NUMBER, 12),
        ]
        # -- each placeholder takes its position from the master's --
        for placeholder in layout.placeholders:
            base = master.placeholders.get(placeholder.placeholder_format.type)
            assert base is not None
            assert (placeholder.left, placeholder.width) == (base.left, base.width)
            assert placeholder._element.spPr.xfrm is None

    def and_it_gives_the_footer_placeholders_the_masters_text(self):
        layout = Presentation().slide_masters[0].slide_layouts.add_slide_layout("Quote")

        slide_number = layout.placeholders.get(idx=12)

        assert slide_number is not None
        assert slide_number._element.xpath(".//a:fld[@type='slidenum']")

    def but_it_adds_no_placeholders_when_blank(self):
        layout = Presentation().slide_masters[0].slide_layouts.add_slide_layout("Empty", blank=True)

        assert len(layout.shapes) == 0

    def it_marks_the_layout_as_user_made_and_kept_when_unused(self):
        layout = Presentation().slide_masters[0].slide_layouts.add_slide_layout("Quote")

        assert layout._element.preserve is True
        assert layout._element.userDrawn is True

    def it_survives_a_save_and_reload_with_a_slide_on_it(self):
        prs = Presentation()
        layout = prs.slide_masters[0].slide_layouts.add_slide_layout("Quote")
        layout.shapes.add_placeholder(PP_PLACEHOLDER.BODY)
        slide = prs.slides.add_slide(layout)
        title = slide.shapes.title
        assert title is not None
        title.text_frame.text = "To be or not to be"

        reloaded = _reloaded(prs)

        new_layout = reloaded.slide_masters[0].slide_layouts.get_by_name("Quote")
        assert new_layout is not None
        new_slide = reloaded.slides[0]
        assert new_slide.slide_layout == new_layout
        new_title = new_slide.shapes.title
        assert new_title is not None
        assert new_title.text_frame.text == "To be or not to be"
        assert [p.placeholder_format.idx for p in new_slide.placeholders] == [0, 13]
        # -- the slide's placeholders inherit their position from the layout --
        layout_body = new_layout.placeholders.get(idx=13)
        assert layout_body is not None
        assert new_slide.placeholders[13].left == layout_body.left

    def it_gives_every_master_and_layout_a_unique_id_in_the_shared_range(self):
        prs = Presentation()
        layouts = prs.slide_masters[0].slide_layouts
        layouts.add_slide_layout("One")
        layouts.add_slide_layout("Two")
        layouts.duplicate(layouts[0])

        ids = _master_and_layout_ids(_reloaded(prs))

        assert None not in ids
        assert len(ids) == len(set(ids))
        assert all(2147483648 <= id <= 4294967295 for id in ids)


class DescribeSlideLayouts_duplicate:
    """`SlideLayouts.duplicate()`."""

    def it_adds_a_copy_directly_after_the_original(self):
        layouts = Presentation().slide_masters[0].slide_layouts
        original = layouts[1]
        names = [layout.name for layout in layouts]

        copy = layouts.duplicate(original)

        assert copy.name == "1_" + original.name
        assert [layout.name for layout in layouts] == names[:2] + [copy.name] + names[2:]
        assert _placeholder_summary(copy) == _placeholder_summary(original)
        assert copy.part is not original.part

    def and_it_counts_up_the_prefix_until_the_name_is_free(self):
        layouts = Presentation().slide_masters[0].slide_layouts
        original = layouts[0]

        first, second = layouts.duplicate(original), layouts.duplicate(original)

        assert (first.name, second.name) == ("1_Title Slide", "2_Title Slide")

    def but_it_uses_the_name_it_is_given(self):
        layouts = Presentation().slide_masters[0].slide_layouts

        assert layouts.duplicate(layouts[0], "Cover").name == "Cover"

    def it_leaves_the_original_unchanged_when_the_copy_is_edited(self):
        prs = Presentation()
        layouts = prs.slide_masters[0].slide_layouts
        original = layouts[1]
        before = original._element.xml
        copy = layouts.duplicate(original)

        copy.placeholders[0].left = Inches(3)
        copy.shapes.add_placeholder(PP_PLACEHOLDER.PICTURE, Inches(1), Inches(1), Inches(1), Inches(1))

        assert original._element.xml == before
        reloaded = _reloaded(prs).slide_masters[0].slide_layouts
        assert reloaded[1]._element.xml == before

    def but_it_raises_for_a_layout_of_another_master(self):
        other_layout = Presentation().slide_masters[0].slide_layouts[0]
        layouts = Presentation().slide_masters[0].slide_layouts

        with pytest.raises(SlideError):
            layouts.duplicate(other_layout)


class DescribeLayoutShapes_add_placeholder:
    """`LayoutShapes.add_placeholder()`."""

    def it_inherits_position_and_size_from_the_master_when_none_are_given(self):
        prs = Presentation()
        master = prs.slide_masters[0]
        layout = master.slide_layouts.add_slide_layout("L", blank=True)

        picture = layout.shapes.add_placeholder(PP_PLACEHOLDER.PICTURE)

        assert isinstance(picture, LayoutPlaceholder)
        body = master.placeholders.get(PP_PLACEHOLDER.BODY)
        assert body is not None
        assert (picture.left, picture.top, picture.width, picture.height) == (
            body.left,
            body.top,
            body.width,
            body.height,
        )
        assert picture.placeholder_format.type == PP_PLACEHOLDER.PICTURE
        assert picture.name == "Picture Placeholder 1"
        assert picture.has_text_frame

    def and_it_takes_an_explicit_position_and_size(self):
        layout = Presentation().slide_masters[0].slide_layouts.add_slide_layout("L", blank=True)

        chart = layout.shapes.add_placeholder(
            PP_PLACEHOLDER.CHART, Inches(1), Inches(2), Inches(3), Inches(4), name="Graph"
        )

        assert (chart.left, chart.top, chart.width, chart.height) == (
            Inches(1),
            Inches(2),
            Inches(3),
            Inches(4),
        )
        assert chart.name == "Graph"

    @pytest.mark.parametrize(
        ("ph_types", "expected_idxs"),
        [
            ((PP_PLACEHOLDER.TITLE,), [0]),
            ((PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT), [13, 14]),
            ((PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE), [12, 10]),
            ((PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.DATE), [10, 13]),
        ],
    )
    def it_assigns_idx_as_PowerPoint_does(self, ph_types, expected_idxs):
        layout = Presentation().slide_masters[0].slide_layouts.add_slide_layout("L", blank=True)

        added = [layout.shapes.add_placeholder(t) for t in ph_types]

        assert [p.placeholder_format.idx for p in added] == expected_idxs

    def and_it_uses_a_free_idx_it_is_given(self):
        layout = Presentation().slide_masters[0].slide_layouts.add_slide_layout("L", blank=True)

        assert layout.shapes.add_placeholder(PP_PLACEHOLDER.BODY, idx=1).placeholder_format.idx == 1

    @pytest.mark.parametrize(
        ("call", "message"),
        [
            (lambda s: s.add_placeholder(PP_PLACEHOLDER.HEADER), "not a slide-layout"),
            (lambda s: s.add_placeholder(PP_PLACEHOLDER.BODY, Inches(1)), "all of left"),
            (lambda s: s.add_placeholder(PP_PLACEHOLDER.TITLE), "already has a title"),
            (lambda s: s.add_placeholder(PP_PLACEHOLDER.CENTER_TITLE), "already has a title"),
            (lambda s: s.add_placeholder(PP_PLACEHOLDER.BODY, idx=10), "already in use"),
            (lambda s: s.add_placeholder(PP_PLACEHOLDER.BODY, idx=0), "already in use"),
        ],
    )
    def but_it_raises_on_a_request_it_cannot_honor(self, call, message):
        layout = Presentation().slide_masters[0].slide_layouts.add_slide_layout("L")

        with pytest.raises(ShapeError, match=message):
            call(layout.shapes)

    def and_it_raises_when_the_master_has_nothing_to_inherit_from(self):
        prs = Presentation()
        master = prs.slide_masters[0]
        layout = master.slide_layouts.add_slide_layout("L", blank=True)
        body = master.placeholders.get(PP_PLACEHOLDER.BODY)
        assert body is not None
        body._element.getparent().remove(body._element)

        with pytest.raises(ShapeError, match="no BODY placeholder"):
            layout.shapes.add_placeholder(PP_PLACEHOLDER.PICTURE)
        picture = layout.shapes.add_placeholder(
            PP_PLACEHOLDER.PICTURE, Inches(1), Inches(1), Inches(2), Inches(2)
        )
        assert picture.width == Inches(2)


class DescribeImportSlide_master_and_layout_ids:
    """`Slides.import_slide()` of a slide whose layout must be imported with its master."""

    def it_lists_only_the_imported_layout_under_the_imported_master(self):
        source = Presentation(str(CORPUS_DIR / "layouts.pptx"))
        slide = source.slides[1]
        slide.slide_layout.name = "Imported Layout"
        prs = Presentation()

        prs.slides.import_slide(slide)

        reloaded = _reloaded(prs)
        assert len(reloaded.slide_masters) == 2
        assert [layout.name for layout in reloaded.slide_masters[1].slide_layouts] == [
            "Imported Layout"
        ]
        ids = _master_and_layout_ids(reloaded)
        assert None not in ids
        assert len(ids) == len(set(ids))


class DescribePresentationPart_next_master_or_layout_id:
    """`PresentationPart.next_master_or_layout_id()`."""

    def it_is_one_more_than_the_highest_id_in_use(self):
        prs = Presentation()

        assert prs.part.next_master_or_layout_id() == max(_master_and_layout_ids(prs)) + 1

    def and_it_does_not_reuse_an_id_below_a_gap(self):
        prs = Presentation()
        entries = prs.slide_masters[0]._element.sldLayoutIdLst.sldLayoutId_lst
        entries[2].id = 2147490000  # -- leaves a gap below it --

        next_id = prs.part.next_master_or_layout_id()

        assert next_id == 2147490001

    def but_it_finds_a_free_id_below_once_the_top_of_the_range_is_taken(self):
        prs = Presentation()
        entries = prs.slide_masters[0]._element.sldLayoutIdLst.sldLayoutId_lst
        entries[-1].id = 4294967295

        next_id = prs.part.next_master_or_layout_id()

        assert next_id not in _master_and_layout_ids(prs)
        assert 2147483648 <= next_id < 4294967295
