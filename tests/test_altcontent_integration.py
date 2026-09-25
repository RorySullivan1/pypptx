"""Integration tests for shapes PowerPoint wraps in `mc:AlternateContent` (3D models, zooms,
equations): they are listed in `slide.shapes`, and shape-tree operations act on the wrapper."""

from __future__ import annotations

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.oxml.shapes.shared import tree_elm
from pptx.shapes.altcontent import AlternateContentShape

from .unitutil.altcontent import altcontent_pptx
from .unitutil.cxml import element


def _shapes():
    return Presentation(altcontent_pptx()).slides[0].shapes


class DescribeListingAlternateContentShapes:
    """Wrapped shapes in a file laid out the way PowerPoint saves them."""

    def it_lists_the_shapes_the_selection_pane_shows(self):
        shapes = _shapes()

        assert len(shapes) == 5
        assert [shape.name for shape in shapes] == [
            "TextBox 9",
            "3D Model 10",
            "Slide Zoom 11",
            "TextBox 12",
            "Rectangle 13",
        ]
        assert [getattr(shape, "content_kind", None) for shape in shapes] == [
            None,
            "model3d",
            "zoom",
            "equation",
            "unknown",
        ]
        assert all(isinstance(shape, AlternateContentShape) for shape in list(shapes)[1:])

    def it_places_each_at_its_fallback_or_else_its_choice(self):
        _, model, _, _, unknown = _shapes()

        assert (model.left, model.top, model.width, model.height) == (110, 210, 2990, 1990)
        assert model.has_fallback is True
        assert model.shape_type is MSO_SHAPE_TYPE.MODEL_3D
        assert (unknown.left, unknown.top, unknown.width, unknown.height) == (7000, 7000, 500, 600)
        assert unknown.has_fallback is False

    def it_finds_them_by_name_and_id(self):
        shapes = _shapes()

        assert shapes.get_by_name("Slide Zoom 11").content_kind == "zoom"
        assert shapes.get_by_id(13).content_kind == "equation"
        assert "3D Model 10" in shapes
        assert shapes.index(shapes[2]) == 2

    def it_skips_a_wrapper_that_holds_no_shape(self):
        spTree = element("p:spTree/(p:sp,mc:AlternateContent/(mc:Choice/a:ext,mc:Fallback))")

        assert [elm.tag for elm in spTree.iter_shape_elms()] == [qn("p:sp")]


class DescribeAlternateContentShapeTreeOperations:
    """Moving, copying, removing and grouping a wrapped shape act on the whole wrapper."""

    def it_moves_the_wrapper_to_the_front_and_back(self):
        shapes = _shapes()
        zoom = shapes[2]

        shapes.move_shape_to_front(zoom)
        assert shapes.index(zoom) == 4
        shapes.move_shape_to_back(zoom)
        assert shapes.index(zoom) == 0
        assert tree_elm(zoom.element).tag == qn("mc:AlternateContent")
        assert zoom.element.getparent().tag == qn("mc:Fallback")

    def it_removes_the_wrapper(self):
        shapes = _shapes()

        shapes.remove_shape(shapes[1])

        assert len(shapes) == 4
        assert len(shapes._spTree.xpath("./mc:AlternateContent")) == 3

    def it_duplicates_the_wrapper_with_a_new_id_on_choice_and_fallback(self):
        shapes = _shapes()
        model = shapes[1]

        copy = shapes.duplicate_shape(model)

        assert isinstance(copy, AlternateContentShape)
        assert copy.content_kind == "model3d"
        assert copy.element.getparent().tag == qn("mc:Fallback")
        wrapper = tree_elm(copy.element)
        assert wrapper is not tree_elm(model.element)
        assert set(wrapper.xpath(".//p:cNvPr/@id")) == {str(copy.shape_id)}
        assert copy.shape_id not in (11, 12, 13, 14)
        assert len(shapes) == 6

    def it_groups_the_wrapper(self):
        shapes = _shapes()
        equation = shapes[3]

        group = shapes.add_group_shape([equation])

        (grouped,) = group.shapes
        assert grouped.content_kind == "equation"
        assert grouped.is_in_group is True
        assert tree_elm(grouped.element).getparent() is group.element
        assert len(shapes) == 5
