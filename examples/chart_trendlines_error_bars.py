"""Trendlines and error bars on a clustered-column chart.

Each of three series gets a different trendline (linear with its equation and R-squared
displayed, a second-order polynomial projected forwards and backwards, and a moving average)
and a different kind of error bar (fixed amount, percentage, standard deviation).

The deck is reopened after saving and every value is asserted against what was written, so
running this script is a round-trip check and not just a drawing exercise.

Demonstrates: feature-manifest.md "Charts — Trendlines" and "Charts — Error Bars".
Saves to chart_trendlines_error_bars.pptx.
"""

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_ERROR_BAR_DIRECTION,
    XL_ERROR_BAR_INCLUDE,
    XL_ERROR_BAR_TYPE,
    XL_LEGEND_POSITION,
    XL_TRENDLINE_TYPE,
)
from pptx.util import Inches

FILENAME = "chart_trendlines_error_bars.pptx"

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])

chart_data = CategoryChartData()
chart_data.categories = ["Q1", "Q2", "Q3", "Q4", "Q5", "Q6"]
chart_data.add_series("Actual", (12.1, 14.8, 13.2, 17.9, 19.4, 22.6))
chart_data.add_series("Forecast", (11.5, 15.2, 14.1, 16.8, 20.2, 21.1))
chart_data.add_series("Baseline", (10.0, 11.0, 12.0, 13.0, 14.0, 15.0))

chart = slide.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED,
    Inches(0.6), Inches(1.2), Inches(8.8), Inches(5),
    chart_data,
).chart
chart.has_legend = True
chart.legend.position = XL_LEGEND_POSITION.BOTTOM
chart.legend.include_in_layout = False

actual, forecast, baseline = chart.plots[0].series

# --- Trendlines -------------------------------------------------------------------------
# A linear fit, showing the fitted equation and R-squared on the plot.
linear = actual.trendlines.add(XL_TRENDLINE_TYPE.LINEAR)
linear.name = "Actual trend"
linear.display_equation = True
linear.display_r_squared = True

# A quadratic fit projected two periods forward and one back. `order` is the polynomial
# degree; PowerPoint accepts 2-6.
poly = forecast.trendlines.add(XL_TRENDLINE_TYPE.POLYNOMIAL)
poly.order = 2
poly.forward = 2.0
poly.backward = 1.0

# A moving average over two periods -- `period`, not `order`, drives this type.
moving = baseline.trendlines.add(XL_TRENDLINE_TYPE.MOVING_AVERAGE)
moving.period = 2

# --- Error bars -------------------------------------------------------------------------
# A fixed +/-1.5 unit band.
actual_bars = actual.error_bars.add(XL_ERROR_BAR_TYPE.FIXED_VALUE, 1.5)
actual_bars.include = XL_ERROR_BAR_INCLUDE.BOTH
actual_bars.direction = XL_ERROR_BAR_DIRECTION.Y

# A 10% band, drawn without end caps.
forecast_bars = forecast.error_bars.add(XL_ERROR_BAR_TYPE.PERCENT, 10.0)
forecast_bars.has_end_cap = False

# One standard deviation. ST_DEV takes a value; ST_ERROR and CUSTOM do not.
baseline_bars = baseline.error_bars.add(XL_ERROR_BAR_TYPE.ST_DEV, 1.0)
baseline_bars.include = XL_ERROR_BAR_INCLUDE.PLUS_VALUES

prs.save(FILENAME)
print(f"Saved {FILENAME}")

# --- Round trip -------------------------------------------------------------------------
# Read every value back out of the saved file rather than off the objects above, which would
# only prove the setters worked in memory.
reopened = Presentation(FILENAME)
# The layout contributes a title placeholder, so the chart is not simply shapes[0].
frame = next(shape for shape in reopened.slides[0].shapes if shape.has_chart)
r_actual, r_forecast, r_baseline = frame.chart.plots[0].series

assert len(r_actual.trendlines) == 1, len(r_actual.trendlines)
assert r_actual.trendlines[0].trendline_type == XL_TRENDLINE_TYPE.LINEAR
assert r_actual.trendlines[0].name == "Actual trend"
assert r_actual.trendlines[0].display_equation is True
assert r_actual.trendlines[0].display_r_squared is True

assert r_forecast.trendlines[0].trendline_type == XL_TRENDLINE_TYPE.POLYNOMIAL
assert r_forecast.trendlines[0].order == 2, r_forecast.trendlines[0].order
assert r_forecast.trendlines[0].forward == 2.0
assert r_forecast.trendlines[0].backward == 1.0

assert r_baseline.trendlines[0].trendline_type == XL_TRENDLINE_TYPE.MOVING_AVERAGE
assert r_baseline.trendlines[0].period == 2, r_baseline.trendlines[0].period

assert r_actual.error_bars.has_error_bars is True
assert r_actual.error_bars[0].type == XL_ERROR_BAR_TYPE.FIXED_VALUE
assert r_actual.error_bars[0].value == 1.5, r_actual.error_bars[0].value
assert r_actual.error_bars[0].direction == XL_ERROR_BAR_DIRECTION.Y
assert r_actual.error_bars[0].include == XL_ERROR_BAR_INCLUDE.BOTH

assert r_forecast.error_bars[0].type == XL_ERROR_BAR_TYPE.PERCENT
assert r_forecast.error_bars[0].value == 10.0
assert r_forecast.error_bars[0].has_end_cap is False

assert r_baseline.error_bars[0].type == XL_ERROR_BAR_TYPE.ST_DEV
assert r_baseline.error_bars[0].include == XL_ERROR_BAR_INCLUDE.PLUS_VALUES

print("Round trip verified: 3 trendlines, 3 error-bar kinds")
