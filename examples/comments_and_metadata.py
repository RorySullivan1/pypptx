"""Comments, tags, sections, and custom document properties.

Demonstrates adding legacy and modern (threaded) comments to slides, setting key-value tags,
creating presentation sections, and custom properties.
Saves to comments_and_metadata.pptx.
"""

from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

# Create two slides
slide1 = prs.slides.add_slide(prs.slide_layouts[5])
txBox1 = slide1.shapes.add_textbox(Inches(2), Inches(3), Inches(6), Inches(1))
txBox1.text_frame.text = "Slide with comments and tags"
txBox1.text_frame.paragraphs[0].font.size = Pt(24)

slide2 = prs.slides.add_slide(prs.slide_layouts[5])
txBox2 = slide2.shapes.add_textbox(Inches(2), Inches(3), Inches(6), Inches(1))
txBox2.text_frame.text = "Second slide"
txBox2.text_frame.paragraphs[0].font.size = Pt(24)

# Add comments to slide 1
slide1.comments.add("Alice", "AJ", "Please review the layout.", x=100, y=200)
slide1.comments.add("Bob", "BK", "Looks good to me!", x=100, y=400)
print(f"Slide 1 has {len(slide1.comments)} comments")

# Add a modern threaded comment to slide 2, reply to it, and resolve it
thread = slide2.threaded_comments.add("Can we add a chart here?", "Alice", "AJ")
thread.reply("Added one in the next revision.", "Bob", "BK")
thread.resolved = True
for thread in slide2.threaded_comments:
    replies = ", ".join(f"{r.author}: {r.text}" for r in thread.replies)
    print(f"Thread by {thread.author} (resolved={thread.resolved}): {thread.text} [{replies}]")

# Set tags on slide 1 (key-value metadata)
slide1.tags["status"] = "draft"
slide1.tags["reviewer"] = "alice"
print(f"Tags: status={slide1.tags['status']}, reviewer={slide1.tags['reviewer']}")

# Create sections to organize slides
intro = prs.sections.add("Introduction")
intro.add_slide_id(prs.slides[0].slide_id)

body = prs.sections.add("Content")
body.add_slide_id(prs.slides[1].slide_id)
print(f"Sections: {len(prs.sections)}")

# Set custom document properties
prs.custom_properties["Department"] = "Engineering"
prs.custom_properties["Version"] = 2
prs.custom_properties["Approved"] = True
print(f"Department: {prs.custom_properties['Department']}")

prs.save("comments_and_metadata.pptx")
print("Saved comments_and_metadata.pptx")

# --- Round trip ---
# Reopen the saved deck and assert every piece of metadata survived. Comments, tags, sections
# and custom properties each live in a different part of the package -- comments in their own
# part, tags in a related tags part, sections in the presentation part's extension list, and
# custom properties in docProps/custom.xml -- so this is really four round trips in one.
reopened = Presentation("comments_and_metadata.pptx")
r_slide1, r_slide2 = reopened.slides

# Legacy comments, with their authors.
assert len(r_slide1.comments) == 2, len(r_slide1.comments)
assert [c.author for c in r_slide1.comments] == ["Alice", "Bob"]
assert [c.text for c in r_slide1.comments] == [
    "Please review the layout.",
    "Looks good to me!",
]

# The modern threaded comment, its reply, and the resolved flag.
assert len(r_slide2.threaded_comments) == 1, len(r_slide2.threaded_comments)
r_thread = list(r_slide2.threaded_comments)[0]
assert r_thread.author == "Alice"
assert r_thread.text == "Can we add a chart here?"
assert r_thread.resolved is True
assert len(r_thread.replies) == 1
assert r_thread.replies[0].author == "Bob"
assert r_thread.replies[0].text == "Added one in the next revision."

# Slide tags.
assert r_slide1.tags["status"] == "draft"
assert r_slide1.tags["reviewer"] == "alice"
assert "missing" not in r_slide1.tags

# Sections, and the slide each one holds.
assert [s.name for s in reopened.sections] == ["Introduction", "Content"]
assert reopened.sections[0].slide_ids == (r_slide1.slide_id,)
assert reopened.sections[1].slide_ids == (r_slide2.slide_id,)

# Custom document properties keep their Python types, not just their text.
assert reopened.custom_properties["Department"] == "Engineering"
assert reopened.custom_properties["Version"] == 2
assert isinstance(reopened.custom_properties["Version"], int)
assert reopened.custom_properties["Approved"] is True

print("Round trip verified: comments, threads, tags, sections, typed custom properties")
