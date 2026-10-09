"""Add custom slide layouts to a master, then build slides on them.

Adds a "Quote" layout the way PowerPoint's Insert Layout does (a title plus date, footer and
slide-number placeholders, each positioned by the master), gives it a body placeholder that
inherits the master's body position and a picture placeholder placed explicitly, and duplicates
an existing layout under a new name. Then adds a slide on each new layout and saves.

Open the result in PowerPoint and look at View > Slide Master: the new layouts appear among the
originals, and the slides show the layouts' placeholders. "1_Quote" is the name pypptx gives an
unnamed duplicate; compare it with what PowerPoint's own Duplicate Layout produces.

Demonstrates: feature-manifest.md "Slide Masters & Layouts" (custom layouts).
Saves to custom_layouts.pptx.
"""

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.shapes.placeholder import SlidePlaceholder
from pptx.util import Inches

FILENAME = "custom_layouts.pptx"

prs = Presentation()
master = prs.slide_masters[0]
layouts = master.slide_layouts

# --- A new layout: title and footers come with it, positioned by the master ---------------
quote = layouts.add_slide_layout("Quote")
quote.shapes.add_placeholder(PP_PLACEHOLDER.BODY)  # inherits the master body's position
quote.shapes.add_placeholder(
    PP_PLACEHOLDER.PICTURE, Inches(6.5), Inches(1.75), Inches(3), Inches(4), name="Portrait"
)
print(f"added layout {quote.name!r}:")
for placeholder in quote.placeholders:
    fmt = placeholder.placeholder_format
    print(f"  idx {fmt.idx:>2}  {fmt.type.name:<13} {placeholder.name}")

# --- A copy of an existing layout, placed right after it -----------------------------------
two_content = layouts.get_by_name("Two Content")
assert two_content is not None
copy = layouts.duplicate(two_content, "Two Content (wide gap)")
copy.placeholders[1].width = Inches(3.5)  # editing the copy leaves the original as it was
print(f"duplicated {two_content.name!r} as {copy.name!r} at position {layouts.index(copy)}")
default_copy = layouts.duplicate(quote)  # unnamed: PowerPoint's own "1_" naming
print(f"duplicated {quote.name!r} as {default_copy.name!r}")

# --- Slides on the new layouts ---------------------------------------------------------------
slide = prs.slides.add_slide(quote)
assert slide.shapes.title is not None
slide.shapes.title.text_frame.text = "On simplicity"
body = slide.placeholders[13]
assert isinstance(body, SlidePlaceholder)
body.text_frame.text = "Simplicity is the ultimate sophistication."

slide = prs.slides.add_slide(copy)
assert slide.shapes.title is not None
slide.shapes.title.text_frame.text = "Built on the duplicated layout"

prs.save(FILENAME)
print(f"saved {FILENAME}: {len(layouts)} layouts, {len(prs.slides)} slides")
