"""Process many presentations in parallel, one worker process per CPU core.

Each deck is opened, has a placeholder token replaced in every text run, and is saved to an
output folder; a per-deck summary (slides, runs changed, or the error) is printed as each
finishes. The work for one deck is independent of every other, so the decks are spread across
worker processes with `concurrent.futures.ProcessPoolExecutor`. On 4 cores, 16 decks of 100 slides each ran 3.4x
faster than a plain loop (issue #51); small decks gain less, as process start-up dominates.

Use processes, not threads: pypptx objects (like the lxml trees under them) must not be shared
between threads, and a process pool also sidesteps the GIL. Give each worker a *path*, not a
`Presentation` -- open and save the deck inside the worker so nothing large crosses processes.

Usage::

    python batch_process.py                      # builds 8 demo decks, then processes them
    python batch_process.py IN_DIR OUT_DIR       # processes IN_DIR/*.pptx into OUT_DIR
    python batch_process.py IN_DIR OUT_DIR --find ACME --replace "Acme Corp" --workers 4

Writes to batch_output/ when run without arguments.
"""

from __future__ import annotations

import argparse
import os
import time
from concurrent.futures import ProcessPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


@dataclass
class Result:
    """What one worker reports back for one deck; small and picklable."""

    name: str
    slides: int = 0
    runs_changed: int = 0
    seconds: float = 0.0
    error: str | None = None


def process_deck(src: Path, dst: Path, find: str, replace: str) -> Result:
    """Open `src`, replace `find` with `replace` in every text run, and save to `dst`.

    Runs in a worker process. Any error is caught and reported, so one bad file does not stop
    the batch.
    """
    started = time.perf_counter()
    try:
        prs = Presentation(str(src))
        changed = 0
        for slide in prs.slides:
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if find in run.text:
                            run.text = run.text.replace(find, replace)
                            changed += 1
        prs.save(str(dst))
        return Result(src.name, len(prs.slides), changed, time.perf_counter() - started)
    except Exception as exc:  # noqa: BLE001 -- report every failure, keep the batch going
        return Result(src.name, seconds=time.perf_counter() - started, error=repr(exc))


def process_folder(in_dir: Path, out_dir: Path, find: str, replace: str, workers: int | None):
    """Process every `.pptx` in `in_dir` into `out_dir` using a pool of `workers` processes."""
    sources = sorted(in_dir.glob("*.pptx"))
    out_dir.mkdir(parents=True, exist_ok=True)
    started = time.perf_counter()
    failures = 0
    # -- the `with` block waits for every submitted deck before it exits --
    with ProcessPoolExecutor(max_workers=workers) as pool:
        futures = [
            pool.submit(process_deck, src, out_dir / src.name, find, replace) for src in sources
        ]
        for future in as_completed(futures):
            result = future.result()
            if result.error:
                failures += 1
                print(f"  FAILED {result.name}: {result.error}")
            else:
                print(
                    f"  {result.name}: {result.slides} slides, {result.runs_changed} runs"
                    f" changed ({result.seconds:.2f}s)"
                )
    elapsed = time.perf_counter() - started
    print(
        f"Processed {len(sources)} decks ({failures} failed) in {elapsed:.2f}s"
        f" with {workers or os.cpu_count()} workers -> {out_dir}"
    )


def build_demo_decks(folder: Path, count: int = 8) -> None:
    """Write `count` small decks whose text mentions the token ACME, to process in the demo."""
    folder.mkdir(parents=True, exist_ok=True)
    for n in range(count):
        prs = Presentation()
        for slide_idx in range(20):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            for row in range(5):
                box = slide.shapes.add_textbox(
                    Inches(0.5), Inches(0.5 + row), Inches(8), Inches(0.8)
                )
                box.text_frame.text = f"ACME quarterly report {n}, slide {slide_idx}, row {row}"
        prs.save(str(folder / f"report_{n:02d}.pptx"))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("in_dir", nargs="?", type=Path, help="folder of .pptx files")
    parser.add_argument("out_dir", nargs="?", type=Path, default=Path("batch_output"))
    parser.add_argument("--find", default="ACME")
    parser.add_argument("--replace", default="Acme Corp")
    parser.add_argument(
        "--workers", type=int, default=None, help="worker processes (default: one per CPU)"
    )
    args = parser.parse_args()

    in_dir = args.in_dir
    if in_dir is None:
        in_dir = args.out_dir / "demo_input"
        build_demo_decks(in_dir)
        print(f"Built demo decks in {in_dir}")
        args.out_dir = args.out_dir / "processed"
    process_folder(in_dir, args.out_dir, args.find, args.replace, args.workers)


# -- the guard is required: worker processes import this module, and must not re-run main() --
if __name__ == "__main__":
    main()
