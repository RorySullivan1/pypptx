"""Office 2016+ ("chartex") charts, plus a classic chart with a data table.

Adds one slide per chartex type -- waterfall, histogram, box & whisker, treemap, sunburst and
funnel -- then a column chart showing its data table, and reads each chartex chart back.
Saves to modern_charts.pptx.
"""

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

prs = Presentation()


def add_chart_slide(title, chart_type, chart_data):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    return slide.shapes.add_chart(
        chart_type, Inches(1), Inches(1.5), Inches(8), Inches(5), chart_data
    )


# --- Waterfall: one series, one value per category ---
cash = CategoryChartData()
cash.categories = ["Opening", "Sales", "Services", "Costs", "Tax", "Closing"]
cash.add_series("Cash flow", (120, 85, 40, -95, -20, 130))
add_chart_slide("Waterfall", XL_CHART_TYPE.WATERFALL, cash)

# --- Histogram: raw values; PowerPoint bins them ---
scores = CategoryChartData()
scores.add_series("Scores", (52, 61, 64, 68, 70, 71, 73, 75, 78, 81, 84, 88, 93))
add_chart_slide("Histogram", XL_CHART_TYPE.HISTOGRAM, scores)

# --- Box & whisker: one series per box ---
times = CategoryChartData()
times.add_series("Team A", (12, 15, 14, 18, 22, 13, 16))
times.add_series("Team B", (9, 11, 19, 10, 12, 30, 11))
add_chart_slide("Box & whisker", XL_CHART_TYPE.BOX_WHISKER, times)

# --- Treemap and sunburst: sizes over hierarchical categories ---
sales = CategoryChartData()
for region, products in (("North", ("Bikes", "Helmets")), ("South", ("Bikes", "Locks", "Lights"))):
    category = sales.add_category(region)
    for product in products:
        category.add_sub_category(product)
sales.add_series("Units", (340, 120, 280, 90, 60))
add_chart_slide("Treemap", XL_CHART_TYPE.TREEMAP, sales)
add_chart_slide("Sunburst", XL_CHART_TYPE.SUNBURST, sales)

# --- Funnel: one series, stages in order ---
pipeline = CategoryChartData()
pipeline.categories = ["Leads", "Qualified", "Proposal", "Won"]
pipeline.add_series("Deals", (500, 240, 90, 35))
add_chart_slide("Funnel", XL_CHART_TYPE.FUNNEL, pipeline)

# --- Classic column chart with a data table under the plot ---
quarters = CategoryChartData()
quarters.categories = ["Q1", "Q2", "Q3", "Q4"]
quarters.add_series("East", (19.2, 21.4, 16.7, 23.1))
quarters.add_series("West", (14.5, 18.3, 22.9, 17.6))
chart = add_chart_slide("Data table", XL_CHART_TYPE.COLUMN_CLUSTERED, quarters).chart
chart.has_data_table = True
chart.data_table.show_keys = True

# --- Read the chartex charts back ---
for slide in prs.slides:
    for shape in slide.shapes:
        if shape.has_chartex:
            chartex = shape.chartex
            series = chartex.series[0]
            print(f"{chartex.chart_type.name:12} {series.name!r}: {series.values}")

prs.save("modern_charts.pptx")
print("Saved modern_charts.pptx")
