"""Minimal presentation with a title slide and core properties.

Creates a single title slide with a title and subtitle, sets document
metadata, and saves to hello_world.pptx.
"""

from pptx import Presentation

prs = Presentation()

# Add a title slide (layout index 0)
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Welcome to pypptx"
slide.placeholders[1].text = "Automate PowerPoint with Python"

# Set core document properties
prs.core_properties.author = "pypptx"
prs.core_properties.title = "Hello World Presentation"

prs.save("hello_world.pptx")
print("Saved hello_world.pptx")
