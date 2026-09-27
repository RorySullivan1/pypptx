"""Walk the master -> layout -> slide placeholder inheritance chain.

A slide placeholder usually carries no position or size of its own: PowerPoint resolves those
by walking up to the layout placeholder with the same `idx`, and from there to the master
placeholder with the same type. This script prints each level for a deck built from the
default template, marks which values a slide inherits rather than states, then overrides one
placeholder's geometry to show the resolution flip.

Read-only with respect to the template -- it builds its own deck and only writes the one
override it demonstrates.

Demonstrates: feature-manifest.md "Slide Masters & Layouts" and "Placeholders".
Saves to custom_slide_master_walk.pptx.
"""

from pptx import Presentation
from pptx.util import Emu, Inches

FILENAME = "custom_slide_master_walk.pptx"


def inches(value):
    """Format an EMU length as inches, or '-' when the value is absent."""
    return "-" if value is None else f"{Emu(value).inches:.2f}in"


def describe(shape):
    return f"{inches(shape.left)} x {inches(shape.top)}  {inches(shape.width)}x{inches(shape.height)}"


prs = Presentation()

# Layout 1 is "Title and Content" in the default template: a title and a body placeholder.
layout = prs.slide_layouts[1]
master = layout.slide_master
slide = prs.slides.add_slide(layout)

print(f"master: {len(master.placeholders)} placeholders, layout: {len(layout.placeholders)}")

# --- The walk ---------------------------------------------------------------------------
for placeholder in slide.placeholders:
    idx = placeholder.placeholder_format.idx
    ph_type = placeholder.placeholder_format.type

    # Layout placeholders are matched to slide placeholders by `idx`; master placeholders are
    # matched to layout placeholders by type, which is why the master lookup differs.
    layout_ph = next(
        (p for p in layout.placeholders if p.placeholder_format.idx == idx), None
    )
    master_ph = next(
        (p for p in master.placeholders if p.placeholder_format.type == ph_type), None
    )

    # `spPr/xfrm` is where a shape states its own geometry. Absent, the value the API returns
    # was resolved from an ancestor -- which is the whole point of this walk.
    states_own_geometry = placeholder._element.spPr.xfrm is not None
    origin = "states its own" if states_own_geometry else "inherits"

    print(f"\nidx={idx} {ph_type} -- {origin}")
    print(f"  slide : {describe(placeholder)}")
    if layout_ph is not None:
        print(f"  layout: {describe(layout_ph)}")
    if master_ph is not None:
        print(f"  master: {describe(master_ph)}")

    # An untouched slide placeholder resolves to exactly its layout ancestor's geometry.
    if not states_own_geometry and layout_ph is not None:
        assert placeholder.left == layout_ph.left, (placeholder.left, layout_ph.left)
        assert placeholder.top == layout_ph.top
        assert placeholder.width == layout_ph.width
        assert placeholder.height == layout_ph.height

# --- Override one placeholder ------------------------------------------------------------
body = slide.placeholders[1]
inherited_left = body.left
layout_body = next(p for p in layout.placeholders if p.placeholder_format.idx == 1)

body.left = Inches(2.5)
assert body.left == Inches(2.5)
# The layout is untouched: the override lives on the slide placeholder alone.
assert layout_body.left != Inches(2.5) or inherited_left == Inches(2.5)
print(f"\nbody idx=1: inherited {inches(inherited_left)} -> overridden {inches(body.left)}")

prs.save(FILENAME)
print(f"Saved {FILENAME}")

# --- Round trip -------------------------------------------------------------------------
reopened = Presentation(FILENAME)
r_slide = reopened.slides[0]
r_layout = r_slide.slide_layout
r_body = r_slide.placeholders[1]
r_title = r_slide.placeholders[0]
r_layout_body = next(p for p in r_layout.placeholders if p.placeholder_format.idx == 1)

# The override survived, and it is stored on the slide rather than pushed up to the layout.
assert r_body.left == Inches(2.5), r_body.left
assert r_body._element.spPr.xfrm is not None
assert r_layout_body.left == layout_body.left, (r_layout_body.left, layout_body.left)

# The title was never touched, so it still resolves upwards.
r_layout_title = next(p for p in r_layout.placeholders if p.placeholder_format.idx == 0)
assert r_title._element.spPr.xfrm is None
assert r_title.left == r_layout_title.left
assert r_title.height == r_layout_title.height

print("Round trip verified: override on the slide, title still inherited")
